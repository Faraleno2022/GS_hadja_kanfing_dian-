"""Liste des révisions payées, exports et ticket compact de paiement."""
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape
from urllib.parse import urlencode

from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, render

from eleves.models import Classe
from eleves.utils_annee import get_annee_active
from utilisateurs.utils import filter_by_user_school, user_school
from .models import Paiement
from .revisions import MONTANT_REVISION, PRECISION_REVISION


def _scope(request):
    qs = filter_by_user_school(Paiement.objects.select_related(
        'eleve__classe__ecole', 'mode_paiement', 'type_paiement',
    ), request.user, 'eleve__classe__ecole')
    annees = list(qs.order_by('-annee_scolaire').values_list('annee_scolaire', flat=True).distinct())
    annee = request.GET.get('annee_scolaire', '').strip()
    ecole = user_school(request.user)
    if not annee:
        annee = (get_annee_active(request, ecole) if ecole else None) or (annees[0] if annees else '')
    qs = qs.filter(statut='VALIDE', frais_revision_inclus=True)
    if annee:
        qs = qs.filter(annee_scolaire=annee)
    classe = request.GET.get('classe', '').strip()
    if classe:
        qs = qs.filter(eleve__classe_id=int(classe)) if classe.isascii() and classe.isdecimal() else qs.none()
    q = request.GET.get('q', '').strip()
    if q:
        qs = qs.filter(Q(eleve__nom__icontains=q) | Q(eleve__prenom__icontains=q) | Q(eleve__matricule__icontains=q))
    classes = filter_by_user_school(Classe.objects.all(), request.user, 'ecole')
    if annee:
        classes = classes.filter(annee_scolaire=annee)
    return qs.order_by('eleve__nom', 'eleve__prenom'), {
        'annees': annees, 'annee': annee, 'classes': classes.order_by('nom'), 'classe': classe, 'q': q,
    }


def _resume_revision(qs, context):
    # Une seule option active par élève et année est garantie en base.
    nombre = qs.count()
    filtres = urlencode({
        key: context[source]
        for key, source in (('annee_scolaire', 'annee'), ('classe', 'classe'), ('q', 'q'))
        if context[source]
    })
    return {
        'nombre': nombre,
        'montant': nombre * int(MONTANT_REVISION),
        'tarif': int(MONTANT_REVISION),
        'annee': context['annee'],
        'filtres': filtres,
    }


def resume_revisions(request):
    """Même périmètre pour la carte, sa liste et son actualisation."""
    qs, context = _scope(request)
    return _resume_revision(qs, context)


@login_required
def liste_revisions(request):
    qs, context = _scope(request)
    resume = _resume_revision(qs, context)
    nombre = resume['nombre']
    query = request.GET.copy()
    query.pop('page', None)
    context.update(titre_page='Élèves ayant payé la révision', page_obj=Paginator(qs, 40).get_page(request.GET.get('page')),
                   nombre=nombre, total_revision=resume['montant'], resume_revision=resume, filtres=query.urlencode())
    return render(request, 'paiements/liste_revisions.html', context)


def _cell(value):
    # Un nom saisi par l'utilisateur ne doit pas être exécuté comme formule Excel.
    return "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value


@login_required
def export_revisions_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    qs, context = _scope(request)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Révisions payées'
    ws.append([f"Révisions payées — {context['annee']}"])
    ws.append([PRECISION_REVISION + ', une fois par année scolaire.'])
    ws.append(['Matricule', 'Élève', 'Classe', 'École', 'Année', 'Date', 'Reçu', 'Versement encaissé (GNF)', 'Réduction révision (GNF)'])
    count = 0
    for paiement in qs:
        count += 1
        ws.append([_cell(v) for v in (
            paiement.eleve.matricule, paiement.eleve.nom_complet, paiement.eleve.classe.nom,
            paiement.eleve.classe.ecole.nom, paiement.annee_scolaire,
            paiement.date_paiement.strftime('%d/%m/%Y'), paiement.numero_recu,
            int(paiement.montant), int(MONTANT_REVISION),
        )])
    ws.append(['TOTAL RÉVISION', '', '', '', '', '', '', '', count * int(MONTANT_REVISION)])
    for cell in ws[3]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='167D8D')
        cell.alignment = Alignment(wrap_text=True)
    for column, width in zip('ABCDEFGHI', [18, 30, 24, 28, 16, 14, 18, 26, 28]):
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=4, min_col=8):
        for cell in row:
            cell.number_format = '#,##0'
    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = f'A3:I{3 + count}'
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="revisions_payees.xlsx"'
    return response


@login_required
def export_revisions_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
    qs, context = _scope(request)
    styles = getSampleStyleSheet()
    styles['BodyText'].fontSize = 9
    def p(value):
        return Paragraph(escape(str(value)), styles['BodyText'])
    data = [[p(value) for value in ['Matricule', 'Élève', 'Classe', 'École', 'Date / reçu', 'Révision (GNF)']]]
    count = 0
    for paiement in qs:
        count += 1
        data.append([p(v) for v in (
            paiement.eleve.matricule, paiement.eleve.nom_complet, paiement.eleve.classe.nom,
            paiement.eleve.classe.ecole.nom,
            paiement.date_paiement.strftime('%d/%m/%Y') + ' / ' + paiement.numero_recu, '20 000',
        )])
    buffer = BytesIO()
    table = Table(data, colWidths=[85, 145, 120, 125, 150, 105], repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e8f6f8')),
        ('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
    ]))
    SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=24, rightMargin=24).build([
        Paragraph('Élèves ayant payé la révision', styles['Title']),
        p(context['annee']), p(PRECISION_REVISION + ', une fois par année scolaire.'), Spacer(1, 14), table,
        Spacer(1, 14), p(f"{count} élève(s) — Total déduit : {count * 20000:,} GNF".replace(',', ' ')),
    ])
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="revisions_payees.pdf"'
    return response


@login_required
def ticket_paiement_pdf(request, paiement_id):
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from .payment_engine import situation_echeancier
    from .allocation import get_payment_allocation
    paiement = get_object_or_404(filter_by_user_school(
        Paiement.objects.select_related('eleve__classe__ecole', 'type_paiement', 'mode_paiement'),
        request.user, 'eleve__classe__ecole',
    ), pk=paiement_id, statut='VALIDE')
    style = ParagraphStyle('ticket', fontName='Helvetica', fontSize=9, leading=12, spaceAfter=5)
    heading = ParagraphStyle('titre', parent=style, fontName='Helvetica-Bold', fontSize=12, alignment=1, leading=15)
    def p(value, bold=False):
        return Paragraph(escape(str(value)), heading if bold else style)
    def money(value):
        return f'{int(value):,} GNF'.replace(',', ' ')
    school = paiement.eleve.classe.ecole
    story = [p(school.nom, True), p('TICKET DE PAIEMENT', True),
             p(paiement.numero_recu), p(paiement.date_paiement.strftime('%d/%m/%Y')),
             p(paiement.eleve.nom_complet), p(f'{paiement.eleve.matricule} — {paiement.eleve.classe.nom}'),
             p(f'Année : {paiement.annee_scolaire}'), p(paiement.type_paiement.nom),
             p(f'Mode : {paiement.mode_paiement.nom}'), p('Encaissé : ' + money(paiement.montant), True)]
    if paiement.frais_revision_inclus:
        story.extend([p('Frais de révision payés'), p(PRECISION_REVISION)])
    echeancier = paiement.echeancier_annuel
    if echeancier:
        situation = situation_echeancier(echeancier)
        allocation = get_payment_allocation(paiement, echeancier)
        story.append(Spacer(1, 5))
        for cle, label in [('inscription', echeancier.libelle_frais_admission), ('tranche_1', 'Tranche 1'), ('tranche_2', 'Tranche 2'), ('tranche_3', 'Tranche 3')]:
            story.append(p(label + ' : ' + money(allocation.get(cle, Decimal('0')))))
        story.extend([p('Total annuel : ' + money(situation['total_du'])),
                      p('Remises annuelles : ' + money(situation['total_remises'])),
                      p('Reste à payer : ' + money(situation['solde_restant']), True)])
    buffer = BytesIO()
    # Adapter la longueur du rouleau au contenu, sans page blanche à découper.
    longueur = 12 * mm + 14 + sum(
        item.wrap(70 * mm - 12, 10000)[1] + item.getSpaceBefore() + item.getSpaceAfter()
        for item in story
    )
    SimpleDocTemplate(buffer, pagesize=(80 * mm, max(100 * mm, longueur)), topMargin=6 * mm, bottomMargin=6 * mm,
                      leftMargin=5 * mm, rightMargin=5 * mm).build(story)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="ticket_{paiement.numero_recu}.pdf"'
    return response
