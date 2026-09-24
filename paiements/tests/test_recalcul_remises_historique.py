from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib import admin
from django.contrib.auth import get_user_model
from django.test import RequestFactory, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from administration.audit import mettre_en_corbeille, restaurer_element
from administration.corbeille import restaurer
from administration.models import ElementCorbeille
from eleves.models import Classe, Ecole, Eleve, GrilleTarifaire
from paiements.admin import PaiementAdmin
from paiements.models import EcheancierPaiement, ModePaiement, Paiement, PaiementRemise, RemiseReduction, TypePaiement
from paiements.payment_engine import recalculer_remises_paiement
from paiements.recalcul_remises import memoriser_regle_remise
from paiements.services import synchroniser_echeancier_apres_changement_paiement
from paiements.tests.support import MIDDLEWARE_SANS_LICENCE


@override_settings(MIDDLEWARE=MIDDLEWARE_SANS_LICENCE)
class RecalculRemisesHistoriqueTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser('recalcul-historique', 'test@example.com', 'test')
        self.client.force_login(self.user)
        self.ecole = Ecole.objects.create(nom='Ecole calculs', adresse='Conakry', telephone='620000001', directeur='Direction')
        self.classe = Classe.objects.create(ecole=self.ecole, nom='Petite section A', niveau='PETITE_SECTION', annee_scolaire='2026-2027')
        self.eleve = Eleve.objects.create(matricule='HIST-001', prenom='Aminata', nom='Diallo', sexe='F', classe=self.classe)
        self.type = TypePaiement.objects.create(nom='Inscription et scolarité annuelle')
        self.mode = ModePaiement.objects.create(nom='Especes historique')
        self.jour = date(2026, 8, 20)
        self.echeancier = self.creer_echeancier(self.eleve)
        self.premier = self.paiement(100000, self.jour)
        self.suivant = self.paiement(100000, self.jour + timedelta(days=1))
        self.remise = RemiseReduction.objects.create(
            nom='Moitié scolarité', type_remise='POURCENTAGE', valeur=50, motif='AUTRE',
            date_debut=self.jour, date_fin=date(2027, 8, 31),
        )
        self.ligne = PaiementRemise.objects.create(
            paiement=self.suivant, remise=self.remise, montant_remise=50000,
            applique_tranche_1=True, montant_tranche_1=50000,
            base_calcul='ECHEANCE', montant_base=100000,
            regle_calcul=memoriser_regle_remise(self.remise, 'ECHEANCE', [1]),
        )
        self.recalculer()

    def creer_echeancier(self, eleve, annee='2026-2027'):
        return EcheancierPaiement.objects.create(
            eleve=eleve, annee_scolaire=annee, frais_inscription_du=100000,
            tranche_1_due=400000, tranche_2_due=0, tranche_3_due=0,
            date_echeance_inscription=date(2026, 10, 1), date_echeance_tranche_1=date(2027, 1, 15),
            date_echeance_tranche_2=date(2027, 3, 15), date_echeance_tranche_3=date(2027, 5, 15),
        )

    def paiement(self, montant, jour, *, statut='VALIDE', eleve=None, annee='2026-2027'):
        return Paiement.objects.create(
            eleve=eleve or self.eleve, montant=montant, date_paiement=jour,
            type_paiement=self.type, mode_paiement=self.mode, statut=statut, annee_scolaire=annee,
        )

    def recalculer(self):
        synchroniser_echeancier_apres_changement_paiement(self.eleve.pk, '2026-2027')
        self.echeancier.refresh_from_db()
        self.ligne.refresh_from_db()

    def modifier(self, paiement, **overrides):
        payload = dict(type_paiement=paiement.type_paiement_id, mode_paiement=self.mode.pk,
                       montant=str(paiement.montant), date_paiement=paiement.date_paiement.isoformat(),
                       motif_modification='Correction vérifiée')
        payload.update(overrides)
        return self.client.post(reverse('paiements:modifier_paiement', args=[paiement.pk]), payload)

    def test_suppression_et_restauration_recalculent_le_versement_suivant(self):
        self.assertEqual(self.echeancier.solde_restant, 250000)
        entree = mettre_en_corbeille(self.premier)
        self.recalculer()
        self.assertEqual(self.ligne.montant_remise, 0)
        self.assertEqual(self.echeancier.solde_restant, 400000)
        restaurer_element(entree)
        self.recalculer()
        self.assertEqual(self.ligne.montant_remise, 50000)
        self.assertEqual(self.echeancier.solde_restant, 250000)

    def test_correction_montant_recalcule_toutes_les_remises(self):
        self.assertEqual(self.modifier(self.premier, montant='50000').status_code, 302)
        self.recalculer()
        self.assertEqual(self.ligne.montant_remise, 25000)
        self.assertEqual(self.ligne.montant_tranche_1, 25000)
        self.assertEqual(self.echeancier.solde_restant, 325000)

    def test_correction_date_rejoue_ordre_chronologique(self):
        self.assertEqual(self.modifier(self.premier, date_paiement='2026-08-22').status_code, 302)
        self.ligne.refresh_from_db()
        self.assertEqual(self.ligne.montant_remise, 0)

    def test_validation_du_premier_recalcule_remise_en_attente(self):
        from paiements.views import _valider_paiement_impl
        self.premier.statut = 'EN_ATTENTE'
        self.premier.save()
        self.suivant.statut = 'EN_ATTENTE'
        self.suivant.save()
        self.recalculer()
        self.assertEqual(self.ligne.montant_remise, 0)
        _valider_paiement_impl(self.premier, self.user)
        self.ligne.refresh_from_db()
        self.assertEqual(self.ligne.montant_remise, 50000)
        self.echeancier.refresh_from_db()
        self.assertEqual(self.echeancier.solde_restant, 400000)

    def test_remise_ancienne_sans_regle_conservee(self):
        self.ligne.regle_calcul = {}
        self.ligne.save()
        mettre_en_corbeille(self.premier)
        self.recalculer()
        self.assertEqual(self.ligne.montant_remise, 50000)

    def test_changement_catalogue_ne_change_pas_le_taux_accorde(self):
        self.remise.valeur = 80
        self.remise.save()
        self.recalculer()
        self.assertEqual(self.ligne.montant_remise, 50000)

    def test_restauration_preserve_la_regle_json(self):
        regle = self.ligne.regle_calcul.copy()
        entree = mettre_en_corbeille(self.suivant)
        paiement, _ = restaurer_element(entree)
        self.ligne = paiement.remises.get()
        self.assertEqual(self.ligne.regle_calcul, regle)
        mettre_en_corbeille(self.premier)
        self.ligne.refresh_from_db()
        self.assertEqual(self.ligne.montant_remise, 0)

    def test_erreur_recalcul_annule_la_suppression(self):
        identifiant = self.premier.pk
        with patch('paiements.recalcul_remises.recalculer_remises_echeancier', side_effect=ValueError('Échec simulé')):
            with self.assertRaises(ValueError):
                mettre_en_corbeille(self.premier)
        self.assertTrue(Paiement.objects.filter(pk=identifiant).exists())

    def test_admin_modification_et_suppression_recalculent(self):
        request = RequestFactory().post('/admin/')
        request.user = self.user
        model_admin = PaiementAdmin(Paiement, admin.site)
        self.premier.montant = Decimal('50000')
        model_admin.save_model(request, self.premier, SimpleNamespace(changed_data=['montant']), True)
        self.ligne.refresh_from_db()
        self.assertEqual(self.ligne.montant_remise, 25000)
        with patch.object(model_admin, 'message_user'):
            model_admin.delete_queryset(request, Paiement.objects.filter(pk=self.premier.pk))
        self.ligne.refresh_from_db()
        self.assertEqual(self.ligne.montant_remise, 0)
        entree = ElementCorbeille.objects.filter(model_label='paiements.Paiement', type_operation=ElementCorbeille.SUPPRESSION).latest('pk')
        restaurer(entree, user=self.user)
        self.ligne.refresh_from_db()
        self.assertEqual(self.ligne.montant_remise, 25000)

    def test_autre_annee_ne_modifie_pas_la_remise(self):
        self.creer_echeancier(self.eleve, annee='2025-2026')
        self.paiement(300000, date(2025, 10, 1), annee='2025-2026')
        self.recalculer()
        self.assertEqual(self.ligne.montant_remise, 50000)

    def grille(self, niveau, inscription=100000, tranche=400000, annee='2026-2027'):
        return GrilleTarifaire.objects.create(
            ecole=self.ecole, niveau=niveau, annee_scolaire=annee,
            frais_inscription=inscription, frais_reinscription=50000,
            tranche_1=tranche, tranche_2=0, tranche_3=0,
        )

    def test_inscription_reinscription_recalcule_frais_et_remise(self):
        self.grille(self.classe.niveau)
        type_reinsc = TypePaiement.objects.create(nom='Réinscription et scolarité annuelle')
        self.assertEqual(self.modifier(self.premier, type_paiement=type_reinsc.pk).status_code, 302)
        self.recalculer()
        self.assertEqual(self.echeancier.frais_inscription_du, 50000)
        self.assertEqual(self.echeancier.nature_frais, 'REINSCRIPTION')
        self.assertEqual(self.echeancier.solde_restant, 200000)

    def test_transfert_meme_niveau_conserve_montants(self):
        self.grille(self.classe.niveau)
        cible = Classe.objects.create(ecole=self.ecole, nom='Petite section B', niveau=self.classe.niveau, annee_scolaire='2026-2027')
        self.eleve.classe = cible
        self.eleve.save()
        self.recalculer()
        self.assertEqual(self.ligne.montant_remise, 50000)
        self.assertEqual(self.echeancier.solde_restant, 250000)

    def test_transfert_tarif_different_recalcule_remise_conserve_cash(self):
        self.ligne.base_calcul = 'TRANCHE'
        self.ligne.regle_calcul = memoriser_regle_remise(self.remise, 'TRANCHE', [1])
        self.ligne.save()
        self.grille('MOYENNE_SECTION', inscription=150000, tranche=600000)
        cible = Classe.objects.create(ecole=self.ecole, nom='Moyenne section B', niveau='MOYENNE_SECTION', annee_scolaire='2026-2027')
        self.eleve.classe = cible
        self.eleve.save()
        self.recalculer()
        self.assertEqual(self.ligne.montant_remise, 300000)
        self.assertEqual(self.echeancier.total_du, 750000)
        self.assertEqual(self.echeancier.solde_restant, 250000)
        self.premier.refresh_from_db()
        self.suivant.refresh_from_db()
        self.assertEqual(self.premier.montant + self.suivant.montant, 200000)

    def test_transfert_nouvelle_annee_isole_historique(self):
        self.grille(self.classe.niveau, annee='2027-2028')
        cible = Classe.objects.create(ecole=self.ecole, nom='Petite section B', niveau=self.classe.niveau, annee_scolaire='2027-2028')
        self.eleve.classe = cible
        self.eleve.save()
        self.ligne.refresh_from_db()
        self.assertEqual(self.ligne.montant_remise, 50000)
        nouveau = self.eleve.echeanciers.get(annee_scolaire='2027-2028')
        self.assertEqual(nouveau.total_paye, 0)
        self.assertEqual(nouveau.total_remises_valides, 0)

    def test_remise_deduite_recalcul_repetable_sans_double_deduction(self):
        self.suivant.montant = Decimal('50000')
        self.suivant.save()
        self.ligne.deduite_du_paiement = True
        self.ligne.save()
        recalculer_remises_paiement(self.suivant)
        self.recalculer()
        self.recalculer()
        self.suivant.refresh_from_db()
        self.assertEqual(self.ligne.montant_remise, 50000)
        self.assertEqual(self.suivant.montant, 50000)

    def test_remise_deduite_transfert_conserve_encaissement_et_reste_stable(self):
        self.test_remise_deduite_recalcul_repetable_sans_double_deduction()
        self.grille('MOYENNE_SECTION', inscription=150000)
        cible = Classe.objects.create(ecole=self.ecole, nom='Moyenne section B', niveau='MOYENNE_SECTION', annee_scolaire='2026-2027')
        self.eleve.classe = cible
        self.eleve.save()
        self.recalculer()
        self.recalculer()
        self.suivant.refresh_from_db()
        self.assertEqual(self.ligne.montant_remise, 25000)
        self.assertEqual(self.suivant.montant, 50000)
        self.assertEqual(self.modifier(self.suivant, observations='Précision administrative').status_code, 302)
        self.recalculer()
        self.suivant.refresh_from_db()
        self.assertEqual(self.ligne.montant_remise, 25000)
        self.assertEqual(self.suivant.montant, 50000)

    def test_cumul_remises_plafonne_et_inscription_exclue(self):
        self.ligne.base_calcul = 'TRANCHE'
        self.ligne.regle_calcul = memoriser_regle_remise(self.remise, 'TRANCHE', [1])
        self.ligne.save()
        autre = RemiseReduction.objects.create(nom='Autre', type_remise='POURCENTAGE', valeur=80, motif='AUTRE', date_debut=self.jour, date_fin=self.jour)
        paiement = self.paiement(10000, self.jour + timedelta(days=2))
        ligne = PaiementRemise.objects.create(paiement=paiement, remise=autre, montant_remise=0, applique_tranche_1=True, regle_calcul=memoriser_regle_remise(autre, 'TRANCHE', [1]))
        self.recalculer()
        ligne.refresh_from_db()
        self.assertEqual(self.ligne.montant_remise + ligne.montant_remise, 400000)
        self.assertEqual(self.echeancier.frais_inscription_du, 100000)


    def test_formulaire_memorise_regle_et_taux_exporte(self):
        self.suivant.statut = 'EN_ATTENTE'
        self.suivant.save()
        response = self.client.post(reverse('paiements:appliquer_remise', args=[self.suivant.pk]), {
            'montant_original': '100000', 'pourcentage_scolarite': '50',
            'tranches': ['1'], 'base_calcul': 'ECHEANCE', 'motif': 'GESTE_COMMERCIAL',
        })
        self.assertEqual(response.status_code, 302)
        ligne = self.suivant.remises.get()
        self.assertEqual(ligne.regle_calcul['base'], 'ECHEANCE')
        self.assertEqual(ligne.regle_calcul['tranches'], [1])
        self.assertEqual(Decimal(ligne.regle_calcul['valeur']), 50)
        self.suivant.statut = 'VALIDE'
        self.suivant.save()
        ligne.remise.valeur = 80
        ligne.remise.save()
        synchroniser_echeancier_apres_changement_paiement(self.eleve.pk, '2026-2027')
        from paiements.views_tranches import _tranche_export_rows
        ligne_export = _tranche_export_rows(self.classe, '2026-2027')[0]
        self.assertEqual(ligne_export['discount_rates'], (Decimal('50'),))
        self.assertEqual(ligne_export['discount'], 50000)

    def test_recu_et_carnet_apres_correction(self):
        self.modifier(self.premier, montant='50000')
        from paiements.carnet_paiement import construire_carnet_paiement_pdf
        pdf, resume = construire_carnet_paiement_pdf(self.suivant, self.echeancier)
        self.assertTrue(pdf.startswith(b'%PDF'))
        self.assertEqual(resume['total_remises'], 25000)
        response = self.client.get(reverse('paiements:generer_recu_pdf', args=[self.suivant.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))

    def test_erreur_recalcul_annule_restauration_admin(self):
        from administration.corbeille import enregistrer_suppression
        entree = enregistrer_suppression(self.premier, user=self.user)
        identifiant = self.premier.pk
        self.premier.delete()
        self.recalculer()
        with patch('paiements.recalcul_remises.recalculer_remises_echeancier', side_effect=ValueError('Échec simulé')):
            with self.assertRaises(ValueError):
                restaurer(entree, user=self.user)
        self.assertFalse(Paiement.objects.filter(pk=identifiant).exists())
        entree.refresh_from_db()
        self.assertFalse(entree.restaure)



@override_settings(MIDDLEWARE=MIDDLEWARE_SANS_LICENCE)
class RecapitulatifVersementsMultiplesTests(TestCase):
    creer_echeancier = RecalculRemisesHistoriqueTests.creer_echeancier
    paiement = RecalculRemisesHistoriqueTests.paiement
    recalculer = RecalculRemisesHistoriqueTests.recalculer
    # Exécuter uniquement les tests définis ici avec un jeu de données autonome.
    def setUp(self):
        RecalculRemisesHistoriqueTests.setUp(self)
        self.ligne.delete()
        self.eleve2 = Eleve.objects.create(matricule='HIST-002', prenom='Moussa', nom='Diallo', sexe='M', classe=self.classe)
        self.creer_echeancier(self.eleve2)
        self.paiement(100000, self.jour, eleve=self.eleve2)
        self.paiement(50000, self.jour, statut='EN_ATTENTE')
        self.paiement(50000, self.jour, statut='REJETE')

    def test_recap_et_recherche_comptent_du_une_fois_par_eleve(self):
        for query, count, due in [('', 2, 1000000), ('Diallo', 2, 1000000), ('HIST-001', 1, 500000)]:
            with self.subTest(query=query):
                response = self.client.get(reverse('paiements:liste_paiements'), {'annee': '2026-2027', 'q': query})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.context['totaux_du']['eleves_count'], count)
                self.assertEqual(response.context['totaux_du']['du_global_net'], due)
                self.assertEqual(sum(row['du_global_net'] for row in response.context['totaux_du_detail_classes']), due)

    def test_export_excel_recap_concorde_avec_ecran(self):
        response = self.client.get(reverse('paiements:export_recap_par_classe_excel'), {'annee': '2026-2027'})
        self.assertEqual(response.status_code, 200)
        rows = list(load_workbook(BytesIO(response.content), data_only=True).active.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][2], 2)
        self.assertEqual(rows[0][7], 1000000)

    def test_transfert_recap_compte_eleve_dans_sa_seule_classe_cible(self):
        cible = Classe.objects.create(ecole=self.ecole, nom='Petite section B', niveau=self.classe.niveau, annee_scolaire='2026-2027')
        self.eleve.classe = cible
        self.eleve.save()
        response = self.client.get(reverse('paiements:liste_paiements'), {'annee': '2026-2027'})
        rows = response.context['totaux_du_detail_classes']
        self.assertEqual(len(rows), 2)
        self.assertEqual([row['eleves_count'] for row in rows], [1, 1])
        self.assertEqual(sum(row['du_global_net'] for row in rows), 1000000)

    def test_rapport_recouvrement_ne_multiplie_pas_le_du(self):
        from paiements.rapports_professionnels import collect_recovery_data
        request = RequestFactory().get('/paiements/', {'classe_id': self.classe.pk, 'annee_scolaire': '2026-2027'})
        request.user = self.user
        data = collect_recovery_data(request)
        self.assertEqual(data['schedule_count'], 2)
        self.assertEqual(data['total_due'], 1000000)
        self.assertEqual(data['total_cash'], 300000)

    def test_export_tranches_ne_multiplie_pas_le_du(self):
        from paiements.views_tranches import _tranche_export_rows
        rows = _tranche_export_rows(self.classe, '2026-2027')
        self.assertEqual(len(rows), 2)
        self.assertEqual(sum(row['total_due'] for row in rows), 1000000)
        self.assertEqual(sum(row['total_paid'] for row in rows), 300000)
