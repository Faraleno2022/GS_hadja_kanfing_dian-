from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader

from administration.audit import mettre_en_corbeille, restaurer_element
from administration.corbeille import annuler_modification
from administration.models import ElementCorbeille
from eleves.models import Classe, Ecole, Eleve, GrilleTarifaire
from paiements.models import EcheancierPaiement, ModePaiement, Paiement, PaiementRemise, TypePaiement
from paiements.services import synchroniser_echeancier_apres_changement_paiement, reconcilier_transfert_classe
from paiements.tests.support import MIDDLEWARE_SANS_LICENCE
from paiements.views import _valider_paiement_impl
from paiements.views_tranches import _tranche_export_rows


@override_settings(MIDDLEWARE=MIDDLEWARE_SANS_LICENCE)
class RevisionTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser('revision', 'revision@example.com', 'test')
        self.client.force_login(self.user)
        self.ecole = Ecole.objects.create(nom='École de démonstration', adresse='Conakry', telephone='620000001', directeur='Direction')
        self.classe = Classe.objects.create(ecole=self.ecole, nom='Petite section A', niveau='PETITE_SECTION', annee_scolaire='2026-2027')
        self.eleve = Eleve.objects.create(matricule='REV-001', prenom='Aminata', nom='Diallo', sexe='F', classe=self.classe)
        self.type = TypePaiement.objects.create(nom='Inscription et scolarité annuelle')
        self.mode = ModePaiement.objects.create(nom='Espèces révision')
        self.echeancier = self.creer_echeancier()

    def creer_echeancier(self, annee='2026-2027'):
        return EcheancierPaiement.objects.create(
            eleve=self.eleve, annee_scolaire=annee, frais_inscription_du=0,
            tranche_1_due=100000, tranche_2_due=100000, tranche_3_due=100000,
            date_echeance_inscription=date(2026, 9, 1), date_echeance_tranche_1=date(2026, 10, 1),
            date_echeance_tranche_2=date(2027, 1, 1), date_echeance_tranche_3=date(2027, 4, 1),
        )

    def paiement(self, revision=True, montant=100000, valide=True, annee='2026-2027'):
        p = Paiement.objects.create(eleve=self.eleve, type_paiement=self.type, mode_paiement=self.mode,
                                    montant=montant, date_paiement=date(2026, 9, 8),
                                    annee_scolaire=annee, frais_revision_inclus=revision)
        synchroniser_echeancier_apres_changement_paiement(self.eleve.pk, annee)
        return _valider_paiement_impl(p, self.user) if valide else p

    def donnees(self, p=None, **extra):
        data = dict(eleve=self.eleve.pk, type_paiement=self.type.pk, mode_paiement=self.mode.pk,
                    montant='100000', date_paiement='2026-09-08', frais_revision_inclus='on',
                    confirmation_paiement_superieur='on', motif_modification='Correction de l’option')
        if p:
            data['montant'] = str(p.montant)
        data.update(extra)
        return data

    def solde(self):
        self.echeancier.refresh_from_db()
        return self.echeancier.solde_restant

    def test_exemple_confirme_et_tranches(self):
        p = self.paiement()
        self.assertEqual(p.montant, Decimal('100000'))
        self.assertEqual(self.solde(), Decimal('180000'))
        self.assertEqual(self.echeancier.total_paye, Decimal('100000'))
        remise = p.remises.get(origine_revision=True)
        self.assertEqual(remise.montant_remise, 20000)
        self.assertFalse(remise.deduite_du_paiement)
        self.assertEqual(sum((remise.montant_tranche_1, remise.montant_tranche_2, remise.montant_tranche_3)), 20000)
        rows = _tranche_export_rows(self.classe, '2026-2027')
        self.assertEqual(rows[0]['balance'], 180000)
        self.assertIn('Révision', rows[0]['precision'])

    def test_sans_option_aucune_reduction(self):
        p = self.paiement(revision=False)
        self.assertEqual(self.solde(), 200000)
        self.assertFalse(p.remises.exists())

    def test_en_attente_pas_de_reduction_effective_ni_de_liste(self):
        p = self.paiement(valide=False)
        self.assertEqual(self.solde(), 300000)
        response = self.client.get(reverse('paiements:liste_revisions'))
        self.assertEqual(response.context['nombre'], 0)
        _valider_paiement_impl(p, self.user)
        self.assertEqual(self.solde(), 180000)
        response = self.client.get(reverse('paiements:liste_revisions'))
        self.assertEqual(response.context['nombre'], 1)

    @patch('paiements.views.send_enrollment_confirmation')
    @patch('paiements.views.send_payment_receipt')
    def test_creation_formulaire_option_enregistree(self, *_):
        response = self.client.post(reverse('paiements:ajouter_paiement'), self.donnees())
        self.assertEqual(response.status_code, 302)
        p = Paiement.objects.get(eleve=self.eleve)
        self.assertTrue(p.frais_revision_inclus)
        self.assertEqual(p.remises.get().montant_remise, 20000)
        _valider_paiement_impl(p, self.user)
        self.assertEqual(self.solde(), 180000)

    @patch('paiements.views.send_payment_receipt')
    def test_creation_excessive_annule_tout(self, send):
        avant = list(EcheancierPaiement.objects.values())
        response = self.client.post(reverse('paiements:ajouter_paiement'), self.donnees(montant='290000'))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Paiement.objects.exists())
        self.assertFalse(PaiementRemise.objects.exists())
        self.assertEqual(list(EcheancierPaiement.objects.values()), avant)
        send.assert_not_called()

    def test_second_versement_ne_deduit_pas_une_deuxieme_fois(self):
        self.paiement()
        self.paiement(revision=False, montant=50000)
        self.assertEqual(self.solde(), 130000)
        with self.assertRaises(ValidationError):
            self.paiement()
        self.assertEqual(Paiement.objects.count(), 2)
        self.assertEqual(PaiementRemise.objects.count(), 1)

    def test_contrainte_base_protege_meme_un_update_direct(self):
        self.paiement()
        p = self.paiement(revision=False, montant=50000)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Paiement.objects.filter(pk=p.pk).update(frais_revision_inclus=True)

    def test_modification_retrait_et_restauration_option(self):
        p = self.paiement()
        response = self.client.post(reverse('paiements:modifier_paiement', args=[p.pk]), self.donnees(p, frais_revision_inclus=''))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.solde(), 200000)
        self.assertFalse(p.remises.exists())
        element = ElementCorbeille.objects.filter(type_operation=ElementCorbeille.MODIFICATION, objet_id=p.pk).latest('pk')
        annuler_modification(element)
        p.refresh_from_db()
        self.assertTrue(p.frais_revision_inclus)
        self.assertEqual(self.solde(), 180000)
        self.assertEqual(p.remises.count(), 1)

    def test_modification_excessive_conserve_montants_et_option(self):
        p = self.paiement()
        avant = list(PaiementRemise.objects.values())
        response = self.client.post(reverse('paiements:modifier_paiement', args=[p.pk]), self.donnees(p, montant='290000'))
        self.assertEqual(response.status_code, 200)
        p.refresh_from_db()
        self.assertEqual(p.montant, 100000)
        self.assertTrue(p.frais_revision_inclus)
        self.assertEqual(list(PaiementRemise.objects.values()), avant)
        self.assertEqual(self.solde(), 180000)

    def test_suppression_et_restauration(self):
        p = self.paiement()
        entree = mettre_en_corbeille(p)
        self.assertEqual(self.solde(), 300000)
        restaurer_element(entree)
        self.assertEqual(self.solde(), 180000)
        self.assertEqual(PaiementRemise.objects.count(), 1)

    def test_restauration_refuse_un_doublon_sans_modification(self):
        ancien = self.paiement()
        entree = mettre_en_corbeille(ancien)
        nouveau = self.paiement()
        with self.assertRaises(Exception):
            restaurer_element(entree)
        self.assertEqual(list(Paiement.objects.values_list('pk', flat=True)), [nouveau.pk])
        self.assertEqual(self.solde(), 180000)

    def test_nouvelle_annee_option_independante(self):
        self.paiement()
        prochain = self.creer_echeancier('2027-2028')
        self.paiement(annee='2027-2028')
        self.assertEqual(self.solde(), 180000)
        self.assertEqual(prochain.solde_restant, 180000)

    def test_transfert_meme_annee_garde_revision_et_encaissement(self):
        self.paiement()
        cible = Classe.objects.create(ecole=self.ecole, nom='Grande section B', niveau='GRANDE_SECTION', annee_scolaire='2026-2027')
        GrilleTarifaire.objects.create(ecole=self.ecole, niveau=cible.niveau, annee_scolaire='2026-2027', frais_inscription=0,
                                      frais_reinscription=0, tranche_1=200000, tranche_2=100000, tranche_3=100000)
        self.eleve.classe = cible
        self.eleve.save()
        reconcilier_transfert_classe(self.eleve, self.classe, cible)
        self.assertEqual(self.solde(), 280000)
        self.assertEqual(PaiementRemise.objects.get().montant_remise, 20000)
        self.assertEqual(Paiement.objects.get().montant, 100000)

    def test_remise_manuelle_preserve_revision_et_refuse_exces(self):
        p = self.paiement(valide=False)
        url = reverse('paiements:appliquer_remise', args=[p.pk])
        data = dict(montant_original='100000', pourcentage_scolarite='10', tranches=['1'], base_calcul='TRANCHE', motif='GESTE_COMMERCIAL')
        self.assertEqual(self.client.post(url, data).status_code, 302)
        self.assertEqual(p.remises.count(), 2)
        self.assertEqual(p.remises.get(origine_revision=True).montant_remise, 20000)
        data.update(pourcentage_scolarite='100', tranches=['1', '2', '3'])
        avant = list(p.remises.values())
        self.assertEqual(self.client.post(url, data).status_code, 200)
        self.assertEqual(list(p.remises.values()), avant)
        _valider_paiement_impl(p, self.user)
        self.assertEqual(self.solde(), 170000)

    def test_exports_recus_ticket_carnet_et_liste(self):
        p = self.paiement()
        params = {'annee_scolaire': '2026-2027'}
        for name, args in [('export_revisions_pdf', []), ('generer_recu_pdf', [p.pk]), ('ticket_paiement_pdf', [p.pk]),
                           ('generer_carnet_paiement_pdf', [p.pk]), ('export_paiements_filtres_pdf', []), ('export_tranches_par_classe_pdf', [])]:
            with self.subTest(name=name):
                response = self.client.get(reverse('paiements:' + name, args=args), params)
                self.assertEqual(response.status_code, 200)
                pdf = PdfReader(BytesIO(response.content))
                texte = '\n'.join(page.extract_text() for page in pdf.pages)
                self.assertIn('Révision', texte)
                if name in ('generer_recu_pdf', 'ticket_paiement_pdf', 'generer_carnet_paiement_pdf'):
                    self.assertIn('180 000', texte)
                if name == 'ticket_paiement_pdf':
                    self.assertEqual(len(pdf.pages), 1)
                # Documents fictifs pour le contrôle visuel, dans le dossier de travail.
                dossier = Path(__file__).resolve().parents[2] / 'tmp/pdfs/revision'
                dossier.mkdir(parents=True, exist_ok=True)
                (dossier / (name + '.pdf')).write_bytes(response.content)
        for name in ('export_revisions_excel', 'export_paiements_filtres_excel', 'export_tranches_par_classe_excel'):
            with self.subTest(name=name):
                response = self.client.get(reverse('paiements:' + name), params)
                self.assertEqual(response.status_code, 200)
                wb = load_workbook(BytesIO(response.content))
                texte = ' '.join(str(c.value or '') for ws in wb for row in ws for c in row)
                self.assertIn('Révision', texte)

    def test_filtre_classe_ne_melange_pas_les_eleves(self):
        self.paiement()
        response = self.client.get(reverse('paiements:liste_revisions'), {'classe': '999999'})
        self.assertEqual(response.context['nombre'], 0)


    def test_ajout_option_sur_paiement_deja_valide(self):
        p = self.paiement(revision=False)
        response = self.client.post(reverse('paiements:modifier_paiement', args=[p.pk]), self.donnees(p))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.solde(), 180000)
        p.refresh_from_db()
        self.assertEqual(p.montant, 100000)
        self.assertTrue(p.frais_revision_inclus)

    def test_limite_exacte_solde_sans_surplus(self):
        self.paiement(montant=280000)
        self.assertEqual(self.solde(), 0)
        self.assertEqual(self.echeancier.total_paye, 280000)

    def test_validation_refuse_un_paiement_devenu_excessif(self):
        p = self.paiement(valide=False, montant=200000)
        self.paiement(revision=False, montant=100000)
        with self.assertRaises(ValidationError):
            _valider_paiement_impl(p, self.user)
        p.refresh_from_db()
        self.assertEqual(p.statut, 'EN_ATTENTE')
        self.assertEqual(p.montant, 200000)
        self.assertEqual(self.solde(), 200000)

    def test_recu_public_precise_revision_et_solde(self):
        from paiements.recu_public import generer_token_recu
        p = self.paiement()
        self.client.logout()
        response = self.client.get(reverse('paiements:recu_public_pdf', args=[p.pk]), {'token': generer_token_recu(p.pk)})
        self.assertEqual(response.status_code, 200)
        texte = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(response.content)).pages)
        self.assertIn('Révision', texte)
        self.assertIn('180 000', texte)

    def test_liste_et_exports_respectent_ecole_du_compte(self):
        from utilisateurs.models import Profil
        p = self.paiement()
        user = get_user_model().objects.create_user('revision-autre-ecole', password='test')
        autre = Ecole.objects.create(nom='Autre école', adresse='Conakry', telephone='620000011', directeur='Direction')
        Profil.objects.update_or_create(user=user, defaults={'ecole': autre, 'role': 'COMPTABLE', 'peut_consulter_rapports': True})
        self.client.force_login(user)
        response = self.client.get(reverse('paiements:liste_revisions'))
        self.assertEqual(response.context['nombre'], 0)
        for name in ('export_revisions_excel', 'export_revisions_pdf'):
            response = self.client.get(reverse('paiements:' + name))
            self.assertEqual(response.status_code, 200)
            if name.endswith('pdf'):
                texte = ' '.join(page.extract_text() for page in PdfReader(BytesIO(response.content)).pages)
            else:
                wb = load_workbook(BytesIO(response.content))
                texte = ' '.join(str(c.value or '') for ws in wb for row in ws for c in row)
            self.assertNotIn(self.eleve.matricule, texte)
        self.assertEqual(self.client.get(reverse('paiements:ticket_paiement_pdf', args=[p.pk])).status_code, 404)

    def test_scolarite_inferieure_a_revision_refusee_atomiquement(self):
        self.echeancier.tranche_1_due = 10000
        self.echeancier.tranche_2_due = 0
        self.echeancier.tranche_3_due = 0
        self.echeancier.save()
        response = self.client.post(reverse('paiements:ajouter_paiement'), self.donnees(montant='5000'))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Paiement.objects.exists())
        self.assertEqual(self.solde(), 10000)

    def test_rapports_distinguent_revision_et_encaissement(self):
        from django.test import RequestFactory
        from rapports.utils import remises_par_categorie
        from paiements.export_modes_encaissement import collect_modes_encaissement_data, build_modes_encaissement_workbook, build_modes_encaissement_pdf
        self.paiement()
        categories = remises_par_categorie(Paiement.objects.filter(statut='VALIDE'))
        self.assertEqual(categories['Frais de révision (réduction de scolarité)'], 20000)
        request = RequestFactory().get('/', {'du': '2026-09-01', 'au': '2026-09-30'})
        request.user = self.user
        data = collect_modes_encaissement_data(request)
        self.assertEqual(data['total_amount'], 100000)
        self.assertEqual(data['revision_count'], 1)
        workbook = build_modes_encaissement_workbook(data)
        texte = ' '.join(str(c.value or '') for ws in workbook for row in ws for c in row)
        self.assertIn('Révision', texte)
        self.assertIn('20 000', texte)
        pdf = build_modes_encaissement_pdf(data)
        texte = ' '.join(page.extract_text() for page in PdfReader(pdf).pages)
        self.assertIn('Révision', texte)

    def test_case_revision_grisee_quand_deja_payee(self):
        from paiements.forms import PaiementForm, PaiementModificationForm
        self.assertFalse(PaiementForm(initial={'eleve': self.eleve}).fields['frais_revision_inclus'].disabled)
        avec = self.paiement()
        autre = self.paiement(revision=False, montant=50000)
        champ = PaiementForm(initial={'eleve': self.eleve}).fields['frais_revision_inclus']
        self.assertTrue(champ.disabled)
        self.assertIn(avec.numero_recu, champ.help_text)
        self.assertTrue(PaiementModificationForm(instance=autre).fields['frais_revision_inclus'].disabled)
        # Le paiement qui porte la révision peut toujours la retirer.
        self.assertFalse(PaiementModificationForm(instance=avec).fields['frais_revision_inclus'].disabled)
        response = self.client.get(reverse('paiements:ajax_eleve_info'), {'matricule': self.eleve.matricule})
        self.assertTrue(response.json()['revision_deja_payee'])
        self.assertIn(avec.numero_recu, response.json()['revision_message'])

    @patch('paiements.views.send_enrollment_confirmation')
    @patch('paiements.views.send_payment_receipt')
    def test_case_cochee_forcee_ignoree_si_deja_payee(self, *_):
        self.paiement()
        response = self.client.post(reverse('paiements:ajouter_paiement'), self.donnees(montant='50000'))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Paiement.objects.filter(frais_revision_inclus=True).count(), 1)
        self.assertEqual(PaiementRemise.objects.filter(origine_revision=True).count(), 1)
