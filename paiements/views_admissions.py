"""Listes et exports des élèves inscrits ou réinscrits."""

from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import Http404, HttpResponse
from django.shortcuts import render

from eleves.models import Classe
from eleves.utils_annee import get_annee_active
from utilisateurs.utils import filter_by_user_school, user_is_superadmin, user_school

from .models import EcheancierPaiement


ADMISSION_LABELS = {
    EcheancierPaiement.NATURE_INSCRIPTION: "Élèves inscrits",
    EcheancierPaiement.NATURE_REINSCRIPTION: "Élèves réinscrits",
}


def _nature_or_404(nature):
    nature = (nature or "").upper()
    if nature not in ADMISSION_LABELS:
        raise Http404("Type d'admission inconnu.")
    return nature


def _admission_scope(request, nature):
    nature = _nature_or_404(nature)
    qs = EcheancierPaiement.objects.select_related(
        "eleve",
        "eleve__classe",
        "eleve__classe__ecole",
        "eleve__responsable_principal",
    ).filter(nature_frais=nature)
    qs = filter_by_user_school(qs, request.user, "eleve__classe__ecole")

    selected_year = (
        request.GET.get("annee_scolaire") or request.GET.get("annee") or ""
    ).strip()
    ecole = None if user_is_superadmin(request.user) else user_school(request.user)
    if not selected_year and ecole:
        selected_year = get_annee_active(request, ecole) or ""
    if not selected_year:
        selected_year = (
            qs.order_by("-annee_scolaire")
            .values_list("annee_scolaire", flat=True)
            .first()
            or ""
        )
    if selected_year:
        qs = qs.filter(annee_scolaire=selected_year)

    selected_class = (request.GET.get("classe") or request.GET.get("classe_id") or "").strip()
    if selected_class.isdigit():
        qs = qs.filter(eleve__classe_id=int(selected_class))
    else:
        selected_class = ""

    query = (request.GET.get("q") or "").strip()
    if query:
        qs = qs.filter(
            Q(eleve__matricule__icontains=query)
            | Q(eleve__nom__icontains=query)
            | Q(eleve__prenom__icontains=query)
            | Q(eleve__classe__nom__icontains=query)
            | Q(eleve__responsable_principal__nom__icontains=query)
            | Q(eleve__responsable_principal__telephone__icontains=query)
        )

    qs = qs.order_by("-eleve__date_creation", "-eleve_id")
    classes = Classe.objects.select_related("ecole")
    classes = filter_by_user_school(classes, request.user, "ecole")
    if selected_year:
        classes = classes.filter(annee_scolaire=selected_year)
    classes = classes.order_by("ecole__nom", "niveau", "nom")
    years = (
        filter_by_user_school(
            EcheancierPaiement.objects.all(), request.user, "eleve__classe__ecole"
        )
        .order_by("-annee_scolaire")
        .values_list("annee_scolaire", flat=True)
        .distinct()
    )
    return nature, qs, classes, years, selected_year, selected_class, query


@login_required
def liste_admissions(request, nature):
    nature, qs, classes, years, selected_year, selected_class, query = _admission_scope(
        request, nature
    )
    page_obj = Paginator(qs, 40).get_page(request.GET.get("page"))
    return render(
        request,
        "paiements/liste_admissions.html",
        {
            "titre_page": ADMISSION_LABELS[nature],
            "nature": nature,
            "page_obj": page_obj,
            "classes": classes,
            "annees": years,
            "selected_year": selected_year,
            "selected_class": selected_class,
            "q": query,
        },
    )


def _student_rows(qs):
    for index, schedule in enumerate(qs, start=1):
        student = schedule.eleve
        parent = student.responsable_principal
        yield [
            index,
            student.matricule or "",
            student.nom or "",
            student.prenom or "",
            student.get_sexe_display() if student.sexe else "",
            student.classe.nom if student.classe_id else "",
            student.classe.ecole.nom if student.classe_id else "",
            schedule.annee_scolaire,
            student.date_inscription.strftime("%d/%m/%Y") if student.date_inscription else "",
            parent.nom_complet if parent else "",
            parent.telephone if parent else "",
        ]


@login_required
def export_admissions_excel(request, nature):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    nature, qs, _classes, _years, selected_year, _selected_class, _query = _admission_scope(
        request, nature
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Inscrits" if nature == EcheancierPaiement.NATURE_INSCRIPTION else "Réinscrits"
    title = f"{ADMISSION_LABELS[nature]} — {selected_year or 'Toutes années'}"
    ws.merge_cells("A1:K1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=15, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].alignment = Alignment(horizontal="center")
    headers = [
        "N°", "Matricule", "Nom", "Prénom", "Sexe", "Classe", "École",
        "Année scolaire", "Date inscription", "Responsable", "Téléphone",
    ]
    ws.append([])
    ws.append(headers)
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center")
    for row in _student_rows(qs):
        ws.append(row)
    widths = [7, 17, 20, 20, 10, 22, 28, 17, 18, 26, 18]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:K{max(3, ws.max_row)}"

    output = BytesIO()
    wb.save(output)
    slug = "inscrits" if nature == EcheancierPaiement.NATURE_INSCRIPTION else "reinscrits"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="eleves_{slug}_{selected_year or "tous"}.xlsx"'
    return response


@login_required
def export_admissions_pdf(request, nature):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from ecole_moderne.pdf_utils import draw_logo_watermark
    from ecole_moderne.branding import get_reportlab_palette

    nature, qs, _classes, _years, selected_year, _selected_class, _query = _admission_scope(
        request, nature
    )
    schedules = list(qs)
    school_ids = {
        item.eleve.classe.ecole_id
        for item in schedules
        if item.eleve.classe_id and item.eleve.classe.ecole_id
    }
    school = schedules[0].eleve.classe.ecole if schedules and len(school_ids) == 1 else None
    palette = get_reportlab_palette(school)

    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "AdmissionTitle", parent=styles["Title"], alignment=TA_CENTER,
        fontSize=16, leading=20, textColor=palette['primary'],
    )
    story = [
        Paragraph(school.nom if school else "MySchoolGN", title_style),
        Paragraph(
            f"{ADMISSION_LABELS[nature]} — Année scolaire {selected_year or 'toutes'}",
            title_style,
        ),
        Spacer(1, 6 * mm),
    ]
    headers = ["N°", "Matricule", "Nom et prénom", "Sexe", "Classe", "École", "Responsable", "Téléphone"]
    data = [headers]
    for row in _student_rows(schedules):
        data.append([row[0], row[1], f"{row[2]} {row[3]}".strip(), row[4], row[5], row[6], row[9], row[10]])
    if len(data) == 1:
        data.append(["—", "", "Aucun élève", "", "", "", "", ""])
    table = Table(data, repeatRows=1, colWidths=[10*mm, 26*mm, 43*mm, 16*mm, 37*mm, 48*mm, 47*mm, 31*mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), palette['header']),
        ("TEXTCOLOR", (0, 0), (-1, 0), palette['header_text']),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.35, palette['border']),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, palette['table_alt']]),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(table)

    def decorate(canvas, _doc):
        width, height = landscape(A4)
        draw_logo_watermark(canvas, width, height, opacity=0.035, ecole=school)
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawRightString(width - 10*mm, 7*mm, f"Page {_doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    slug = "inscrits" if nature == EcheancierPaiement.NATURE_INSCRIPTION else "reinscrits"
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="eleves_{slug}_{selected_year or "tous"}.pdf"'
    return response
