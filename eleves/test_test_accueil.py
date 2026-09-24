from datetime import date, timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from paiements.tests.support import MIDDLEWARE_SANS_LICENCE
from openpyxl import load_workbook

from .models import Classe, Ecole, Eleve, Responsable


@override_settings(
    MIDDLEWARE=MIDDLEWARE_SANS_LICENCE,
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class TestAccueilElevesTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser(
            username="admin-test-accueil",
            email="accueil@example.com",
            password="test-password",
        )
        self.client.force_login(self.user)
        self.school = Ecole.objects.create(
            nom="École test accueil",
            adresse="Conakry",
            telephone="620000901",
            directeur="Direction",
        )
        self.classe = Classe.objects.create(
            ecole=self.school,
            nom="1ère accueil",
            niveau="PRIMAIRE_1",
            annee_scolaire="2026-2027",
        )
        self.parent = Responsable.objects.create(
            nom="Parent",
            prenom="Accueil",
            relation="PERE",
            telephone="620000902",
            adresse="Conakry",
        )
        self.old_student = self._student("ACC-001", "Ancien")
        self.new_student = self._student("ACC-002", "Nouveau")
        Eleve.objects.filter(pk=self.old_student.pk).update(
            date_creation=timezone.now() - timedelta(days=2)
        )

    def _student(self, matricule, prenom):
        return Eleve.objects.create(
            matricule=matricule,
            prenom=prenom,
            nom="Élève",
            sexe="F",
            date_naissance=date(2018, 1, 1),
            classe=self.classe,
            date_inscription=date(2026, 8, 1),
            responsable_principal=self.parent,
        )

    def test_derniers_eleves_sont_affiches_en_premier(self):
        response = self.client.get(reverse("eleves:liste_eleves"))
        self.assertEqual(response.status_code, 200)
        students = list(response.context["page_obj"].object_list)
        self.assertEqual(students[0].pk, self.new_student.pk)

    def test_pointage_test_accueil(self):
        response = self.client.post(
            reverse("eleves:pointer_test_accueil", args=[self.old_student.pk]),
            {"evalue": "1"},
        )
        self.assertEqual(response.status_code, 302)
        self.old_student.refresh_from_db()
        self.assertTrue(self.old_student.test_accueil_evalue)

    def test_exports_evalues_pdf_et_excel(self):
        self.new_student.test_accueil_evalue = True
        self.new_student.save(update_fields=["test_accueil_evalue", "date_modification"])
        excel = self.client.get(
            reverse("eleves:export_test_accueil_excel", args=["evalues"])
        )
        pdf = self.client.get(
            reverse("eleves:export_test_accueil_pdf", args=["evalues"])
        )
        self.assertEqual(excel.status_code, 200)
        self.assertTrue(excel.content.startswith(b"PK"))
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        workbook = load_workbook(BytesIO(excel.content), read_only=True)
        values = {
            cell
            for row in workbook.active.iter_rows(values_only=True)
            for cell in row
        }
        self.assertIn("ACC-002", values)
