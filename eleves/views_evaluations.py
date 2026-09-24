"""Pointage et exports du test d'accueil des élèves."""

from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST

from utilisateurs.utils import filter_by_user_school, user_is_superadmin, user_school

from .models import Eleve
from .utils_annee import get_annee_active


STATUS_MAP = {
    "evalues": (True, "Élèves évalués au test d'accueil"),
    "non-evalues": (False, "Élèves non évalués au test d'accueil"),
}


def _evaluation_status_or_404(status):
    try:
        return STATUS_MAP[status]
    except KeyError as exc:
        raise Http404("Statut d'évaluation inconnu.") from exc


def _evaluation_scope(request, status):
    evaluated, label = _evaluation_status_or_404(status)
    qs = Eleve.objects.select_related(
        "classe", "classe__ecole", "responsable_principal"
    ).filter(test_accueil_evalue=evaluated)
    qs = filter_by_user_school(qs, request.user, "classe__ecole")

    selected_year = (
        request.GET.get("annee_scolaire") or request.GET.get("annee") or ""
    ).strip()
    school = None if user_is_superadmin(request.user) else user_school(request.user)
    if not selected_year and school:
        selected_year = get_annee_active(request, school) or ""
    if selected_year:
        qs = qs.filter(classe__annee_scolaire=selected_year)

    class_id = (request.GET.get("classe_id") or request.GET.get("classe") or "").strip()
    if class_id.isdigit():
        qs = qs.filter(classe_id=int(class_id))
    query = (request.GET.get("q") or request.GET.get("recherche") or "").strip()
    if query:
        qs = qs.filter(
            Q(matricule__icontains=query)
            | Q(nom__icontains=query)
            | Q(prenom__icontains=query)
            | Q(classe__nom__icontains=query)
        )
    return qs.order_by("-date_creation", "-pk"), evaluated, label, selected_year


@login_required
@require_POST
def pointer_test_accueil(request, eleve_id):
    qs = filter_by_user_school(
        Eleve.objects.select_related("classe", "classe__ecole"),
        request.user,
        "classe__ecole",
    )
    student = get_object_or_404(qs, pk=eleve_id)
    value = (request.POST.get("evalue") or "").strip().lower()
    if value not in {"1", "0", "true", "false"}:
        messages.error(request, "Statut du test d'accueil invalide.")
    else:
        student.test_accueil_evalue = value in {"1", "true"}
        student.save(update_fields=["test_accueil_evalue", "date_modification"])
        label = "évalué" if student.test_accueil_evalue else "non évalué"
        messages.success(request, f"{student.nom_complet} est maintenant marqué {label}.")

    target = request.POST.get("next") or ""
    if target and url_has_allowed_host_and_scheme(
        target, allowed_hosts={request.get_host()}, require_https=request.is_secure()
    ):
        return redirect(target)
    return redirect("eleves:liste_eleves")


def _rows(qs):
    for index, student in enumerate(qs, start=1):
        parent = student.responsable_principal
        yield [
            index,
            student.matricule or "",
            student.nom or "",
            student.prenom or "",
            student.get_sexe_display() if student.sexe else "",
            student.classe.nom if student.classe_id else "",
            student.classe.ecole.nom if student.classe_id else "",
            student.classe.annee_scolaire if student.classe_id else "",
            student.date_inscription.strftime("%d/%m/%Y") if student.date_inscription else "",
            parent.nom_complet if parent else "",
            parent.telephone if parent else "",
        ]


@login_required
def export_test_accueil_excel(request, status):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    qs, _evaluated, label, selected_year = _evaluation_scope(request, status)
    wb = Workbook()
    ws = wb.active
    ws.title = "Évalués" if status == "evalues" else "Non évalués"
    ws.merge_cells("A1:K1")
    ws["A1"] = f"{label} — {selected_year or 'Toutes années'}"
    ws["A1"].font = Font(bold=True, size=15, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].alignment = Alignment(horizontal="center")
    headers = [
        "N°", "Matricule", "Nom", "Prénom", "Sexe", "Classe", "École",
        "Année scolaire", "Date inscription", "Responsable", "Téléphone",
    ]
    ws.append([])
    ws.append(headers)
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center")
    for row in _rows(qs):
        ws.append(row)
    widths = [7, 17, 20, 20, 10, 22, 28, 17, 18, 26, 18]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:K{max(3, ws.max_row)}"
    output = BytesIO()
    wb.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="test_accueil_{status}.xlsx"'
    return response


@login_required
def export_test_accueil_pdf(request, status):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from ecole_moderne.pdf_utils import draw_logo_watermark
    from ecole_moderne.branding import get_reportlab_palette

    qs, _evaluated, label, selected_year = _evaluation_scope(request, status)
    students = list(qs)
    school_ids = {item.classe.ecole_id for item in students if item.classe_id}
    school = students[0].classe.ecole if students and len(school_ids) == 1 else None
    palette = get_reportlab_palette(school)
    output = BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=landscape(A4), leftMargin=10*mm, rightMargin=10*mm,
        topMargin=12*mm, bottomMargin=12*mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "EvaluationTitle", parent=styles["Title"], alignment=TA_CENTER,
        fontSize=16, leading=20, textColor=palette['primary'],
    )
    story = [
        Paragraph(school.nom if school else "MySchoolGN", title_style),
        Paragraph(f"{label} — {selected_year or 'Toutes années'}", title_style),
        Spacer(1, 6*mm),
    ]
    data = [["N°", "Matricule", "Nom et prénom", "Sexe", "Classe", "École", "Responsable", "Téléphone"]]
    for row in _rows(students):
        data.append([row[0], row[1], f"{row[2]} {row[3]}".strip(), row[4], row[5], row[6], row[9], row[10]])
    if len(data) == 1:
        data.append(["—", "", "Aucun élève", "", "", "", "", ""])
    table = Table(data, repeatRows=1, colWidths=[10*mm, 26*mm, 43*mm, 16*mm, 37*mm, 48*mm, 47*mm, 31*mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), palette['header']),
        ("TEXTCOLOR", (0, 0), (-1, 0), palette['header_text']),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), .35, palette['border']),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, palette['table_alt']]),
    ]))
    story.append(table)

    def decorate(canvas, _doc):
        width, height = landscape(A4)
        draw_logo_watermark(canvas, width, height, opacity=.035, ecole=school)
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawRightString(width - 10*mm, 7*mm, f"Page {_doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="test_accueil_{status}.pdf"'
    return response
