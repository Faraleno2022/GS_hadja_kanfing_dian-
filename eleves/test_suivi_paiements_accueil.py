from datetime import date
from io import BytesIO

from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from paiements.tests import test_school_filtering as fixtures
from paiements.tests.support import MIDDLEWARE_SANS_LICENCE
from paiements.models import EcheancierPaiement, Paiement
from .models import Classe, Eleve


@override_settings(MIDDLEWARE=MIDDLEWARE_SANS_LICENCE)
class ListesActifsEtAccueilTests(TestCase):
    def setUp(self):
        fixtures.SchoolFilteringTests.setUp(self)
        self.user1.refresh_from_db()
        profil = self.user1.profil
        profil.role = 'ADMIN'
        profil.allowed_menus = ['eleves', 'paiements']
        profil.save()
        self.client.force_login(self.user1)
        self.sans = Eleve.objects.create(
            matricule='SANS-1', prenom='Sans', nom='Paiement', sexe='F', classe=self.classe1,
        )

    def liste(self, **params):
        response = self.client.get(reverse('eleves:actifs_paiements'), params)
        self.assertEqual(response.status_code, 200, response.get('Location'))
        return response

    def paiement(self, eleve=None, **params):
        data = dict(eleve=eleve or self.eleve1, type_paiement=self.type_insc,
                    mode_paiement=self.mode_espece, montant=10000,
                    date_paiement=date(2024, 10, 1), statut='VALIDE')
        data.update(params)
        return Paiement.objects.create(**data)

    def test_repartition_compte_chaque_eleve_une_fois(self):
        self.paiement()
        response = self.liste(paiement='avec')
        rows = list(response.context['page_obj'])
        self.assertEqual([e.pk for e in rows], [self.eleve1.pk])
        self.assertEqual(rows[0].nombre_paiements, 2)
        self.assertEqual(rows[0].montant_paye, 40000)
        self.assertEqual(response.context['stats'], {'total': 2, 'avec': 1, 'sans': 1})
        self.assertEqual(response.context['repartition'][0]['avec'], 1)
        self.assertNotContains(response, 'Ecole B')

    def test_attente_rejete_rembourse_et_annee_passee_ne_comptent_pas(self):
        for statut in ('EN_ATTENTE', 'REJETE', 'REMBOURSE'):
            self.paiement(self.sans, statut=statut)
        self.paiement(self.sans, annee_scolaire='2023-2024')
        rows = list(self.liste(paiement='sans').context['page_obj'])
        self.assertEqual([e.pk for e in rows], [self.sans.pk])
        self.assertEqual(rows[0].nombre_paiements, 0)
        self.assertEqual(rows[0].montant_paye, 0)

    def test_validation_remboursement_suppression_actualisent_les_listes(self):
        p = self.paiement(self.sans, statut='EN_ATTENTE')
        self.assertEqual(self.liste().context['stats']['sans'], 1)
        p.statut = 'VALIDE'
        p.save()
        self.assertEqual(self.liste().context['stats']['avec'], 2)
        p.statut = 'REMBOURSE'
        p.save()
        self.assertEqual(self.liste().context['stats']['sans'], 1)
        self.paiement1.delete()
        self.assertEqual(self.liste().context['stats']['sans'], 2)

    def test_inactifs_et_imports_verrouilles_exclus(self):
        for n, statut in enumerate(('SUSPENDU', 'EXCLU', 'TRANSFERE', 'DIPLOME')):
            Eleve.objects.create(matricule=f'INACTIF-{n}', prenom='Inactif', nom=str(n),
                                 classe=self.classe1, sexe='M', statut=statut)
        Eleve.objects.create(matricule='VERROUILLE', prenom='Import', nom='Verrouillé',
                             classe=self.classe1, sexe='F', import_verrouille=True)
        self.assertEqual(self.liste(paiement='tous').context['stats']['total'], 2)

    def test_classe_recherche_et_transfert(self):
        nouvelle = Classe.objects.create(
            ecole=self.ecole1, nom='C1 B', niveau='PRIMAIRE_1', annee_scolaire='2024-2025',
        )
        self.eleve1.classe = nouvelle
        self.eleve1.save()
        response = self.liste(classe_id=nouvelle.pk, paiement='avec', recherche='Alpha')
        self.assertEqual([e.pk for e in response.context['page_obj']], [self.eleve1.pk])
        self.assertEqual(response.context['stats']['avec'], 1)
        self.assertEqual(self.liste(classe_id=self.classe1.pk).context['stats']['avec'], 0)

    def test_isolation_y_compris_admin_et_compte_sans_ecole(self):
        for name in ('eleves:liste_eleves', 'eleves:actifs_paiements'):
            response = self.client.get(reverse(name))
            self.assertNotContains(response, self.eleve2.matricule)
        self.assertEqual(self.client.get(reverse('eleves:actifs_paiements'), {'classe_id': self.classe2.pk}).status_code, 404)
        profil = self.user1.profil
        profil.ecole = None
        profil.save()
        for name in ('eleves:liste_eleves', 'eleves:actifs_paiements'):
            response = self.client.get(reverse(name))
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.context['page_obj'].paginator.count, 0)

    def test_annee_active_et_pagination_conservent_filtres(self):
        ancienne = Classe.objects.create(ecole=self.ecole1, nom='Ancienne',
            niveau='PRIMAIRE_1', annee_scolaire='2023-2024')
        Eleve.objects.create(matricule='ANCIEN', prenom='Ancien', nom='Élève', sexe='M', classe=ancienne)
        for n in range(32):
            Eleve.objects.create(matricule=f'P-{n}', prenom='Sans', nom=f'Paiement {n}',
                                 sexe='M', classe=self.classe1)
        response = self.liste(paiement='sans', classe_id=self.classe1.pk, recherche='Sans', page=2)
        self.assertEqual(response.context['page_obj'].paginator.count, 33)
        self.assertEqual(len(response.context['page_obj']), 3)
        self.assertIn('classe_id=', response.context['filtres_url'])
        session = self.client.session
        session['annee_scolaire_active'] = '2023-2024'
        session.save()
        self.assertEqual(self.liste(paiement='tous').context['stats']['total'], 1)

    def test_accueil_exclut_reinscription_interface_compteurs_exports_et_post(self):
        self.echeancier1.nature_frais = 'REINSCRIPTION'
        self.echeancier1.save()
        response = self.client.get(reverse('eleves:liste_eleves'))
        self.assertContains(response, 'Non concerné')
        self.assertNotContains(response, f'id="evaluation-accueil-{self.eleve1.pk}"')
        self.assertEqual(response.context['stats']['eleves_non_evalues'], 1)
        url = reverse('eleves:definir_evaluation_accueil', args=[self.eleve1.pk])
        self.assertEqual(self.client.post(url, {'est_evalue': '1'}).status_code, 404)
        self.eleve1.refresh_from_db()
        self.assertFalse(self.eleve1.evaluation_accueil_effectuee)
        response = self.client.get(reverse('eleves:export_evaluation_accueil_excel', args=['non-evalues']))
        values = list(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual([row[0] for row in values[1:]], [self.sans.matricule])

    def test_accueil_inscription_pointage_et_changement_nature_sans_cache(self):
        url = reverse('eleves:definir_evaluation_accueil', args=[self.eleve1.pk])
        self.client.get(reverse('eleves:liste_eleves'))
        self.assertEqual(self.client.post(url, {'est_evalue': '1'}).status_code, 302)
        response = self.client.get(reverse('eleves:liste_eleves'))
        self.assertEqual(response.context['stats']['eleves_evalues'], 1)
        self.echeancier1.nature_frais = 'REINSCRIPTION'
        self.echeancier1.save()
        response = self.client.get(reverse('eleves:liste_eleves'))
        self.assertEqual(response.context['stats']['eleves_evalues'], 0)

    def test_accueil_import_verrouille_et_exports_filtrent_la_classe(self):
        verrouille = Eleve.objects.create(matricule='ACC-LOCK', prenom='Import', nom='Verrouillé',
            sexe='M', classe=self.classe1, import_verrouille=True)
        url = reverse('eleves:definir_evaluation_accueil', args=[verrouille.pk])
        self.assertEqual(self.client.post(url, {'est_evalue': '1'}).status_code, 404)
        for classe_id, expected in ((self.classe1.pk, 2), (self.classe2.pk, 0)):
            response = self.client.get(reverse('eleves:export_evaluation_accueil_excel', args=['non-evalues']),
                                       {'classe_id': classe_id})
            self.assertEqual(len(list(load_workbook(BytesIO(response.content)).active.values)) - 1, expected)

    def test_reinscription_ancienne_ne_remplace_pas_admission_courante(self):
        ancien = EcheancierPaiement.objects.create(
            eleve=self.sans, annee_scolaire='2023-2024', nature_frais='REINSCRIPTION',
            date_echeance_inscription=date(2023, 9, 1), date_echeance_tranche_1=date(2023, 10, 1),
            date_echeance_tranche_2=date(2024, 1, 1), date_echeance_tranche_3=date(2024, 4, 1),
        )
        response = self.client.get(reverse('eleves:liste_eleves'))
        self.assertEqual(response.context['stats']['eleves_non_evalues'], 2)
        self.assertContains(response, f'id="evaluation-accueil-{self.sans.pk}"')

    def test_superadmin_repartition_separe_ecoles_et_paiements(self):
        self.user1.is_superuser = True
        self.user1.save()
        response = self.liste(paiement='avec')
        self.assertEqual(response.context['stats'], {'total': 3, 'avec': 2, 'sans': 1})
        self.assertEqual(len(response.context['repartition']), 2)
        self.assertEqual({e.pk for e in response.context['page_obj']}, {self.eleve1.pk, self.eleve2.pk})

    def test_pointage_partiel_retourne_une_page_complete_et_garde_classe(self):
        response = self.client.get(reverse('eleves:liste_eleves'), {
            'partial': '1', 'classe_id': self.classe1.pk, 'recherche': 'Alpha',
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="next" value="/eleves/liste/?classe_id=')
        self.assertNotContains(response, 'name="next" value="/eleves/liste/?partial=')
