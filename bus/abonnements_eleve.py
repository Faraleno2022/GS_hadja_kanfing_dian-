"""Historique des abonnements (bus et cantine) d'un élève.

Regroupe les abonnements d'un élève en lignes homogènes, propose les
informations à reprendre pour un réabonnement, et produit la liste
(PDF / Excel) ainsi que le carnet d'abonnement PDF.
"""

from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape

from dateutil.relativedelta import relativedelta
from django.db.models import Sum
from django.utils import timezone

from .models import AbonnementBus, AbonnementCantine, GrilleTarifaireBus


MOIS_FR = (
    '', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
    'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre',
)
JOURS_FR = ('Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi', 'Samedi', 'Dimanche')

SERVICES = {'bus': 'Bus', 'cantine': 'Cantine'}

DUREES_CANTINE = {
    AbonnementCantine.Periodicite.JOURNALIER: relativedelta(days=1),
    AbonnementCantine.Periodicite.HEBDOMADAIRE: relativedelta(weeks=1),
    AbonnementCantine.Periodicite.MENSUEL: relativedelta(months=1),
    AbonnementCantine.Periodicite.TRIMESTRIEL: relativedelta(months=3),
    AbonnementCantine.Periodicite.ANNUEL: relativedelta(years=1),
}


def expiration_cantine(periodicite, date_debut):
    """Dernier jour couvert par un abonnement cantine commençant à date_debut."""
    duree = DUREES_CANTINE.get(periodicite, relativedelta(months=1))
    return date_debut + duree - timedelta(days=1)


def montant_gnf(valeur):
    return f"{int(Decimal(str(valeur or 0))):,}".replace(',', ' ') + " GNF"


def _statut(abonnement, aujourd_hui):
    if abonnement.statut == 'SUSPENDU':
        return 'Suspendu'
    if abonnement.date_expiration and abonnement.date_expiration < aujourd_hui:
        return 'Expiré'
    return 'Actif'


def lignes_abonnements(eleve, service='tous'):
    """Tous les abonnements de l'élève, du plus ancien au plus récent."""
    aujourd_hui = timezone.localdate()
    lignes = []

    if service in ('tous', 'bus'):
        for abo in (
            AbonnementBus.objects.filter(eleve=eleve)
            .select_related('grille', 'mode_paiement')
        ):
            zone = abo.grille.zone if abo.grille_id else abo.zone
            detail = abo.get_periodicite_display()
            if zone:
                detail += f" – {zone}"
            if abo.annee_scolaire:
                detail += f" ({abo.annee_scolaire})"
            lignes.append({'service': 'Bus', 'detail': detail, 'abonnement': abo,
                           'piece': abo.numero_recu or abo.reference_externe or '',
                           'mode': getattr(abo.mode_paiement, 'nom', '') or ''})

    if service in ('tous', 'cantine'):
        for abo in AbonnementCantine.objects.filter(eleve=eleve):
            lignes.append({
                'service': 'Cantine',
                'detail': f"{abo.get_periodicite_display()} – {abo.get_type_repas_display()}",
                'abonnement': abo,
                'piece': abo.reference_externe or '',
                'mode': '',
            })

    for ligne in lignes:
        abo = ligne.pop('abonnement')
        expiration = abo.date_expiration
        ligne.update({
            'id': abo.pk,
            'mois': f"{MOIS_FR[abo.date_debut.month]} {abo.date_debut.year}" if abo.date_debut else '',
            'montant': abo.montant or Decimal('0'),
            'date_debut': abo.date_debut,
            'date_expiration': expiration,
            'jour_expiration': JOURS_FR[expiration.weekday()] if expiration else '',
            'jours_restants': (expiration - aujourd_hui).days if expiration else None,
            'statut': _statut(abo, aujourd_hui),
        })
    lignes.sort(key=lambda l: (l['date_debut'] or aujourd_hui, l['service'], l['id']))
    return lignes


def totaux(lignes):
    resultat = {'Bus': Decimal('0'), 'Cantine': Decimal('0')}
    for ligne in lignes:
        resultat[ligne['service']] += Decimal(str(ligne['montant']))
    resultat['total'] = resultat['Bus'] + resultat['Cantine']
    return resultat


# ── Reprise des informations pour un réabonnement ────────────────────────────

def reprise_cantine(eleve):
    """Valeurs à pré-remplir pour un nouvel abonnement cantine de l'élève."""
    responsable = getattr(eleve, 'responsable_principal', None)
    dernier = (
        AbonnementCantine.objects.filter(eleve=eleve)
        .order_by('-date_expiration', '-id')
        .first()
    )
    if dernier is None:
        return {
            'existe': False,
            'nb_abonnements': 0,
            'contact_parent': getattr(responsable, 'telephone', '') or '',
        }
    debut = dernier.date_expiration + timedelta(days=1)
    return {
        'existe': True,
        'nb_abonnements': AbonnementCantine.objects.filter(eleve=eleve).count(),
        'derniere_expiration': dernier.date_expiration.isoformat(),
        'type_repas': dernier.type_repas,
        'periodicite': dernier.periodicite,
        'montant': int(dernier.montant or 0),
        'contact_parent': dernier.contact_parent or getattr(responsable, 'telephone', '') or '',
        'regime_alimentaire': dernier.regime_alimentaire,
        'allergies': dernier.allergies,
        'alerte_avant_jours': dernier.alerte_avant_jours,
        'date_debut': debut.isoformat(),
        'date_expiration': expiration_cantine(dernier.periodicite, debut).isoformat(),
    }


def prochaine_tranche_bus(eleve, grille):
    """Première tranche de la grille qui n'est pas encore soldée par l'élève."""
    payes = dict(
        AbonnementBus.objects.filter(eleve=eleve, grille=grille)
        .values_list('periodicite')
        .annotate(total=Sum('montant'))
    )
    for code in ('T1', 'T2', 'T3'):
        du = grille.montant_pour(code) or 0
        if du > 0 and (payes.get(code) or 0) < du:
            return code
    return ''


def reprise_bus(eleve):
    """Valeurs à pré-remplir pour un nouveau paiement bus de l'élève."""
    dernier = (
        AbonnementBus.objects.filter(eleve=eleve)
        .select_related('grille')
        .order_by('-date_debut', '-id')
        .first()
    )
    if dernier is None:
        return {'existe': False, 'nb_abonnements': 0}
    grille = dernier.grille if dernier.grille_id and dernier.grille.actif else None
    ecole_id = getattr(eleve.classe, 'ecole_id', None)
    if grille is None and dernier.zone and ecole_id:
        grille = (
            GrilleTarifaireBus.objects.filter(ecole_id=ecole_id, zone=dernier.zone, actif=True)
            .order_by('-annee_scolaire')
            .first()
        )
    return {
        'existe': True,
        'nb_abonnements': AbonnementBus.objects.filter(eleve=eleve).count(),
        'dernier_paiement': dernier.date_debut.isoformat() if dernier.date_debut else '',
        'derniere_expiration': dernier.date_expiration.isoformat() if dernier.date_expiration else '',
        'grille_id': grille.pk if grille else '',
        'periodicite': prochaine_tranche_bus(eleve, grille) if grille else '',
        'mode_paiement_id': dernier.mode_paiement_id or '',
        'zone': dernier.zone,
        'point_arret': dernier.point_arret,
        'contact_parent': dernier.contact_parent,
    }


def completer_depuis_dernier_bus(abonnement):
    """Recopie les infos logistiques du précédent abonnement bus si absentes."""
    precedent = (
        AbonnementBus.objects.filter(eleve_id=abonnement.eleve_id)
        .exclude(pk=abonnement.pk)
        .order_by('-date_debut', '-id')
        .first()
    )
    if precedent is None:
        return
    for champ in ('itineraire', 'point_arret', 'contact_parent'):
        if not getattr(abonnement, champ):
            setattr(abonnement, champ, getattr(precedent, champ))


# ── Exports ──────────────────────────────────────────────────────────────────

COLONNES = [
    'Mois', 'Service', 'Détail', 'Montant payé', 'Date début', "Date d'expiration",
    "Jour d'expiration", 'Jours restants', 'Statut', 'Reçu / référence',
]


def _jours_restants_texte(ligne):
    jours = ligne['jours_restants']
    if jours is None:
        return ''
    if jours < 0:
        return f"Expiré depuis {-jours} j"
    return f"{jours} j"


def _date(valeur):
    return valeur.strftime('%d/%m/%Y') if valeur else ''


def construire_excel(eleve, lignes):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    classe = getattr(eleve, 'classe', None)
    ecole = getattr(classe, 'ecole', None)
    wb = Workbook()
    ws = wb.active
    ws.title = "Abonnements"

    ws.append([f"Abonnements de {eleve.nom_complet}"])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([
        f"Matricule : {eleve.matricule}",
        f"Classe : {getattr(classe, 'nom', '')}",
        f"École : {getattr(ecole, 'nom', '')}",
        f"Édité le {timezone.localdate().strftime('%d/%m/%Y')}",
    ])
    ws.append([])

    ws.append(COLONNES)
    entete = ws.max_row
    fond = PatternFill('solid', fgColor='1F4E79')
    bord = Border(*(Side(style='thin', color='BFBFBF'),) * 4)
    for cellule in ws[entete]:
        cellule.font = Font(bold=True, color='FFFFFF')
        cellule.fill = fond
        cellule.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for ligne in lignes:
        ws.append([
            ligne['mois'], ligne['service'], ligne['detail'], int(ligne['montant']),
            ligne['date_debut'], ligne['date_expiration'], ligne['jour_expiration'],
            ligne['jours_restants'] if ligne['jours_restants'] is not None else '',
            ligne['statut'], ligne['piece'],
        ])
        rangee = ws.max_row
        ws.cell(rangee, 4).number_format = '#,##0 "GNF"'
        ws.cell(rangee, 5).number_format = 'DD/MM/YYYY'
        ws.cell(rangee, 6).number_format = 'DD/MM/YYYY'

    somme = totaux(lignes)
    ws.append([])
    for libelle, cle in (('Total bus', 'Bus'), ('Total cantine', 'Cantine'), ('Total général', 'total')):
        ws.append(['', '', libelle, int(somme[cle])])
        ws.cell(ws.max_row, 3).font = Font(bold=True)
        ws.cell(ws.max_row, 4).font = Font(bold=True)
        ws.cell(ws.max_row, 4).number_format = '#,##0 "GNF"'

    for rangee in ws.iter_rows(min_row=entete, max_row=entete + len(lignes)):
        for cellule in rangee:
            cellule.border = bord
    for index, largeur in enumerate((16, 10, 38, 16, 13, 15, 15, 14, 11, 20), start=1):
        ws.column_dimensions[get_column_letter(index)].width = largeur
    ws.freeze_panes = ws.cell(entete + 1, 1)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def construire_pdf(eleve, lignes, carnet=False, service='tous'):
    """Liste des abonnements (carnet=False) ou carnet d'abonnement (carnet=True)."""
    from django.contrib.staticfiles import finders
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from ecole_moderne.branding import get_reportlab_palette
    from ecole_moderne.pdf_utils import draw_logo_watermark

    classe = getattr(eleve, 'classe', None)
    ecole = getattr(classe, 'ecole', None)
    palette = get_reportlab_palette(ecole)
    largeur, hauteur = landscape(A4)
    buffer = BytesIO()
    titre_doc = "CARNET D'ABONNEMENT" if carnet else 'LISTE DES ABONNEMENTS'
    if service in SERVICES:
        titre_doc += f" – {SERVICES[service].upper()}"
    document = SimpleDocTemplate(
        buffer, pagesize=(largeur, hauteur),
        leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=12 * mm,
        title=f"{titre_doc} - {eleve.nom_complet}",
        author=getattr(ecole, 'nom', 'MySchoolGN'),
    )

    styles = getSampleStyleSheet()
    s_titre = ParagraphStyle('AboTitre', parent=styles['Title'], fontName='Helvetica-Bold',
                             fontSize=18, leading=21, textColor=palette['primary'], alignment=TA_LEFT)
    s_sous = ParagraphStyle('AboSous', parent=styles['Normal'], fontSize=9, leading=12,
                            textColor=palette['muted'])
    s_petit = ParagraphStyle('AboPetit', parent=styles['Normal'], fontSize=8.5, leading=11,
                             textColor=palette['text'])
    s_cell = ParagraphStyle('AboCell', parent=styles['Normal'], fontSize=8, leading=9.5,
                            alignment=TA_CENTER)
    s_entete = ParagraphStyle('AboEntete', parent=s_cell, fontName='Helvetica-Bold',
                              textColor=palette['header_text'])

    def texte(valeur):
        return escape(str(valeur or ''))

    logo = None
    try:
        chemin_logo = getattr(getattr(ecole, 'logo', None), 'path', None) or finders.find('logos/logo.jpeg')
        if chemin_logo:
            logo = Image(chemin_logo, width=22 * mm, height=22 * mm, kind='proportional')
    except Exception:
        logo = None
    entete = Table([[logo or '', [
        Paragraph(titre_doc, s_titre),
        Paragraph(
            f"{texte(getattr(ecole, 'nom', 'Établissement scolaire'))}<br/>"
            f"Édité le {timezone.localdate().strftime('%d/%m/%Y')}", s_sous),
    ]]], colWidths=[28 * mm, 249 * mm])
    entete.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                                ('LEFTPADDING', (0, 0), (-1, -1), 0)]))

    dernier_bus = next((l for l in reversed(lignes) if l['service'] == 'Bus'), None)
    derniere_cantine = next((l for l in reversed(lignes) if l['service'] == 'Cantine'), None)
    infos = Table([[
        Paragraph(f"<b>Élève</b><br/>{texte(eleve.nom_complet)}", s_petit),
        Paragraph(f"<b>Matricule</b><br/>{texte(eleve.matricule)}", s_petit),
        Paragraph(f"<b>Classe</b><br/>{texte(getattr(classe, 'nom', ''))}", s_petit),
        Paragraph(f"<b>Bus</b><br/>{texte(dernier_bus['detail'] if dernier_bus else '—')}", s_petit),
        Paragraph(f"<b>Cantine</b><br/>{texte(derniere_cantine['detail'] if derniere_cantine else '—')}", s_petit),
    ]], colWidths=[62 * mm, 35 * mm, 45 * mm, 70 * mm, 65 * mm])
    infos.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), palette['primary_soft']),
        ('BOX', (0, 0), (-1, -1), 0.6, palette['border']),
        ('INNERGRID', (0, 0), (-1, -1), 0.35, palette['table']),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 6), ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))

    somme = totaux(lignes)
    resume = Table([[
        Paragraph(f"<b>Abonnements</b><br/>{len(lignes)}", s_cell),
        Paragraph(f"<b>Total bus</b><br/>{montant_gnf(somme['Bus'])}", s_cell),
        Paragraph(f"<b>Total cantine</b><br/>{montant_gnf(somme['Cantine'])}", s_cell),
        Paragraph(f"<b>Total payé</b><br/>{montant_gnf(somme['total'])}", s_cell),
    ]], colWidths=[69.25 * mm] * 4)
    resume.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F7FAFC')),
        ('BACKGROUND', (3, 0), (3, 0), colors.HexColor('#E8F5E9')),
        ('BOX', (0, 0), (-1, -1), 0.6, palette['border']),
        ('INNERGRID', (0, 0), (-1, -1), 0.35, palette['table']),
        ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))

    colonnes = ['N°'] + COLONNES + (['Visa'] if carnet else [])
    largeurs = [9, 24, 16, 55, 26, 21, 22, 21, 23, 16, 26] + ([18] if carnet else [])
    if not carnet:
        largeurs[3] += 18
    tableau = [[Paragraph(texte(c), s_entete) for c in colonnes]]
    for numero, ligne in enumerate(lignes, start=1):
        rangee = [
            str(numero), ligne['mois'], ligne['service'], ligne['detail'],
            montant_gnf(ligne['montant']), _date(ligne['date_debut']),
            _date(ligne['date_expiration']), ligne['jour_expiration'],
            _jours_restants_texte(ligne), ligne['statut'], ligne['piece'],
        ]
        tableau.append([Paragraph(texte(v), s_cell) for v in rangee] + ([''] if carnet else []))
    if carnet:
        while len(tableau) < 11:
            tableau.append([''] * len(colonnes))
    elif not lignes:
        tableau.append([Paragraph('Aucun abonnement enregistré.', s_cell)] + [''] * (len(colonnes) - 1))

    grille = Table(tableau, colWidths=[l * mm for l in largeurs], repeatRows=1,
                   rowHeights=None if not carnet else [9 * mm] + [8 * mm] * (len(tableau) - 1))
    style = [
        ('BACKGROUND', (0, 0), (-1, 0), palette['header']),
        ('GRID', (0, 0), (-1, -1), 0.5, palette['border']),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, palette['table_alt']]),
        ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 2), ('RIGHTPADDING', (0, 0), (-1, -1), 2),
    ]
    if not lignes and not carnet:
        style.append(('SPAN', (0, 1), (-1, 1)))
    for index, ligne in enumerate(lignes, start=1):
        if ligne['statut'] == 'Expiré':
            style.append(('TEXTCOLOR', (8, index), (9, index), colors.HexColor('#C62828')))
    grille.setStyle(TableStyle(style))

    contenu = [entete, Spacer(1, 3 * mm), infos, Spacer(1, 2.5 * mm), resume,
               Spacer(1, 3.5 * mm), grille]
    if carnet:
        signatures = Table([[
            Paragraph('<b>Signature du parent</b>', s_petit),
            Paragraph('<b>Signature du comptable</b>', s_petit),
            Paragraph('<b>Cachet de l\'école</b>', s_petit),
        ]], colWidths=[92 * mm] * 3, rowHeights=[18 * mm])
        signatures.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'),
                                        ('BOX', (0, 0), (-1, -1), 0.5, palette['border']),
                                        ('INNERGRID', (0, 0), (-1, -1), 0.5, palette['border'])]))
        contenu += [Spacer(1, 4 * mm), signatures]

    def _page(canvas, doc):
        draw_logo_watermark(canvas, largeur, hauteur, ecole=ecole, opacity=0.025, rotate=0, scale=0.75)
        canvas.saveState()
        canvas.setFont('Helvetica', 7)
        canvas.setFillColor(colors.HexColor('#718096'))
        canvas.drawRightString(largeur - 10 * mm, 6 * mm, f"Page {doc.page}")
        canvas.restoreState()

    document.build(contenu, onFirstPage=_page, onLaterPages=_page)
    return buffer.getvalue()
