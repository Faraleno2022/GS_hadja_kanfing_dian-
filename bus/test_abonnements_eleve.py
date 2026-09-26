from datetime import date, timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from eleves.models import Classe, Ecole, Eleve
from paiements.models import ModePaiement
from paiements.tests.support import MIDDLEWARE_SANS_LICENCE
from utilisateurs.models import Profil

from .abonnements_eleve import expiration_cantine, lignes_abonnements
from .models import AbonnementBus, AbonnementCantine, GrilleTarifaireBus


@override_settings(MIDDLEWARE=MIDDLEWARE_SANS_LICENCE)
class AbonnementsEleveTests(TestCase):
    def setUp(self):
        self.ecole = Ecole.objects.create(
            nom='École Abo A', adresse='Conakry', telephone='+224620000201', directeur='Dir A',
        )
        autre_ecole = Ecole.objects.create(
            nom='École Abo B', adresse='Conakry', telephone='+224620000202', directeur='Dir B',
        )
        self.classe = Classe.objects.create(
            ecole=self.ecole, nom='CM1', niveau='PRIMAIRE_4', annee_scolaire='2026-2027',
        )
        self.autre_classe = Classe.objects.create(
            ecole=autre_ecole, nom='CM2', niveau='PRIMAIRE_5', annee_scolaire='2026-2027',
        )
        self.eleve = Eleve.objects.create(
            matricule='ABO-A-001', prenom='Aïssatou', nom='Barry', sexe='F', classe=self.classe,
        )
        self.nouvel_eleve = Eleve.objects.create(
            matricule='ABO-A-002', prenom='Mamadou', nom='Diallo', sexe='M', classe=self.classe,
        )
        self.autre_eleve = Eleve.objects.create(
            matricule='ABO-B-001', prenom='Sékou', nom='Touré', sexe='M', classe=self.autre_classe,
        )
        self.mode = ModePaiement.objects.create(nom='Espèces Abo')
        self.grille = GrilleTarifaireBus.objects.create(
            ecole=self.ecole, zone='Kipé', annee_scolaire='2026-2027',
            tranche_1=150000, tranche_2=125000, tranche_3=125000,
            date_echeance_tranche_3=date(2027, 4, 15),
        )
        self.cantine = AbonnementCantine.objects.create(
            eleve=self.eleve, montant=200000, periodicite='MENSUEL', type_repas='COMPLET',
            date_debut=date(2026, 10, 1), date_expiration=date(2026, 10, 31),
            contact_parent='+224620000299', regime_alimentaire='Sans porc', allergies='Arachides',
        )
        self.bus = AbonnementBus.objects.create(
            eleve=self.eleve, grille=self.grille, periodicite='T1', montant=150000,
            date_debut=date(2026, 9, 20), date_expiration=date(2027, 4, 15),
            mode_paiement=self.mode, point_arret='Carrefour Kipé', contact_parent='+224620000299',
        )
        user = get_user_model().objects.create_user(username='admin_abo_a', password='secret123')
        Profil.objects.update_or_create(user=user, defaults={
            'role': 'ADMIN', 'telephone': '+224620000211', 'ecole': self.ecole,
            'is_validated': True, 'allowed_menus': ['bus'],
        })
        self.client.force_login(user)

    # ── Réabonnement sans ressaisie ─────────────────────────────────────────
    def test_expiration_cantine_couvre_la_periode(self):
        self.assertEqual(expiration_cantine('MENSUEL', date(2026, 11, 1)), date(2026, 11, 30))
        self.assertEqual(expiration_cantine('TRIMESTRIEL', date(2026, 10, 1)), date(2026, 12, 31))
        self.assertEqual(expiration_cantine('HEBDOMADAIRE', date(2026, 10, 5)), date(2026, 10, 11))
        self.assertEqual(expiration_cantine('JOURNALIER', date(2026, 10, 5)), date(2026, 10, 5))

    def test_reprise_cantine_reprend_les_infos_et_enchaine_les_dates(self):
        data = self.client.get(
            reverse('bus:reprise_abonnement_json', args=[self.eleve.pk, 'cantine'])
        ).json()
        self.assertTrue(data['existe'])
        self.assertEqual(data['type_repas'], 'COMPLET')
        self.assertEqual(data['montant'], 200000)
        self.assertEqual(data['regime_alimentaire'], 'Sans porc')
        self.assertEqual(data['allergies'], 'Arachides')
        self.assertEqual(data['contact_parent'], '+224620000299')
        self.assertEqual(data['date_debut'], '2026-11-01')
        self.assertEqual(data['date_expiration'], '2026-11-30')

    def test_reprise_refusee_pour_un_eleve_d_une_autre_ecole(self):
        response = self.client.get(
            reverse('bus:reprise_abonnement_json', args=[self.autre_eleve.pk, 'cantine'])
        )
        self.assertEqual(response.status_code, 404)

    def test_formulaire_cantine_pre_rempli_pour_un_eleve_deja_abonne(self):
        response = self.client.get(reverse('bus:creer_abonnement_cantine'), {'eleve': self.eleve.pk})
        form = response.context['form']
        self.assertEqual(form.initial['regime_alimentaire'], 'Sans porc')
        self.assertEqual(form.initial['date_debut'], '2026-11-01')
        self.assertEqual(form.initial['type_repas'], 'COMPLET')

    def test_cantine_sans_date_expiration_la_calcule(self):
        response = self.client.post(reverse('bus:creer_abonnement_cantine'), {
            'eleve': self.eleve.pk, 'montant': 200000, 'periodicite': 'MENSUEL',
            'type_repas': 'COMPLET', 'date_debut': '2026-11-01', 'date_expiration': '',
            'statut': 'ACTIF', 'alerte_avant_jours': 7,
        })
        self.assertRedirects(response, reverse('bus:liste_abonnements_cantine'))
        nouveau = AbonnementCantine.objects.filter(eleve=self.eleve).latest('id')
        self.assertEqual(nouveau.date_expiration, date(2026, 11, 30))
        self.assertEqual(AbonnementCantine.objects.filter(eleve=self.eleve).count(), 2)

    def test_reprise_bus_propose_la_prochaine_tranche(self):
        data = self.client.get(
            reverse('bus:reprise_abonnement_json', args=[self.eleve.pk, 'bus'])
        ).json()
        self.assertTrue(data['existe'])
        self.assertEqual(data['grille_id'], self.grille.pk)
        self.assertEqual(data['periodicite'], 'T2')
        self.assertEqual(data['mode_paiement_id'], self.mode.pk)

        vierge = self.client.get(
            reverse('bus:reprise_abonnement_json', args=[self.nouvel_eleve.pk, 'bus'])
        ).json()
        self.assertFalse(vierge['existe'])

    def test_formulaire_bus_pre_rempli_et_logistique_reprise(self):
        response = self.client.get(reverse('bus:nouveau'), {'eleve': self.eleve.pk})
        initial = response.context['form'].initial
        self.assertEqual(initial['grille'], self.grille.pk)
        self.assertEqual(initial['periodicite'], 'T2')

        response = self.client.post(reverse('bus:nouveau'), {
            'eleve': self.eleve.pk, 'grille': self.grille.pk, 'periodicite': 'T2',
            'montant': 125000, 'date_debut': '2027-01-10', 'mode_paiement': self.mode.pk,
            'observations': '',
        })
        self.assertRedirects(response, reverse('bus:liste'))
        nouveau = AbonnementBus.objects.get(eleve=self.eleve, periodicite='T2')
        self.assertEqual(nouveau.point_arret, 'Carrefour Kipé')
        self.assertEqual(nouveau.contact_parent, '+224620000299')

    def test_eleves_d_une_classe_avec_indicateurs(self):
        data = self.client.get(reverse('bus:eleves_classe_json', args=[self.classe.pk])).json()
        par_id = {e['id']: e for e in data['eleves']}
        self.assertTrue(par_id[self.eleve.pk]['abonne_bus'])
        self.assertTrue(par_id[self.eleve.pk]['abonne_cantine'])
        self.assertFalse(par_id[self.nouvel_eleve.pk]['abonne_cantine'])
        autre = self.client.get(reverse('bus:eleves_classe_json', args=[self.autre_classe.pk]))
        self.assertEqual(autre.status_code, 404)

    # ── Liste, exports et carnet ────────────────────────────────────────────
    def test_lignes_regroupent_bus_et_cantine_avec_jour_d_expiration(self):
        lignes = lignes_abonnements(self.eleve)
        self.assertEqual([l['service'] for l in lignes], ['Bus', 'Cantine'])
        cantine = lignes[1]
        self.assertEqual(cantine['mois'], 'Octobre 2026')
        self.assertEqual(cantine['jour_expiration'], 'Samedi')  # 31/10/2026
        self.assertEqual(
            cantine['jours_restants'], (date(2026, 10, 31) - timezone.localdate()).days
        )
        self.assertEqual(len(lignes_abonnements(self.eleve, 'cantine')), 1)

    def test_page_abonnements_eleve(self):
        response = self.client.get(reverse('bus:abonnements_eleve', args=[self.eleve.pk]))
        self.assertContains(response, 'Kipé')
        self.assertContains(response, 'Déjeuner + Goûter')
        interdit = self.client.get(reverse('bus:abonnements_eleve', args=[self.autre_eleve.pk]))
        self.assertEqual(interdit.status_code, 404)

    def test_export_excel(self):
        from openpyxl import load_workbook

        response = self.client.get(reverse('bus:abonnements_eleve_excel', args=[self.eleve.pk]))
        self.assertEqual(response.status_code, 200)
        feuille = load_workbook(BytesIO(response.content)).active
        valeurs = [c.value for ligne in feuille.iter_rows() for c in ligne]
        self.assertIn("Jour d'expiration", valeurs)
        self.assertIn('Samedi', valeurs)
        self.assertIn(350000, valeurs)

    def _texte_pdf(self, response):
        from pypdf import PdfReader

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        return '\n'.join(p.extract_text() or '' for p in PdfReader(BytesIO(response.content)).pages)

    def test_liste_pdf_et_carnet(self):
        liste = self._texte_pdf(
            self.client.get(reverse('bus:abonnements_eleve_pdf', args=[self.eleve.pk]))
        )
        self.assertIn('LISTE DES ABONNEMENTS', liste)
        self.assertIn('ABO-A-001', liste)

        carnet = self._texte_pdf(
            self.client.get(reverse('bus:carnet_abonnement_pdf', args=[self.eleve.pk]), {'type': 'cantine'})
        )
        self.assertIn("CARNET D'ABONNEMENT", carnet)
        self.assertIn('Octobre 2026', carnet)
        self.assertIn('31/10/2026', carnet)
        self.assertIn('Samedi', carnet)
        self.assertIn('Signature du parent', carnet)

    def test_carnet_eleve_sans_abonnement(self):
        carnet = self._texte_pdf(
            self.client.get(reverse('bus:carnet_abonnement_pdf', args=[self.nouvel_eleve.pk]))
        )
        self.assertIn("CARNET D'ABONNEMENT", carnet)
