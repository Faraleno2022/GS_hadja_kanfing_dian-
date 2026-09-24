from datetime import timedelta
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from eleves.models import Classe, Ecole, Eleve
from paiements.tests.support import MIDDLEWARE_SANS_LICENCE
from .models_recouvrement import AbonnementInformatique, DepenseCuisine


@override_settings(MIDDLEWARE=MIDDLEWARE_SANS_LICENCE)
class RecouvrementAuditTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('audit-recouvrement')
        self.ecole = Ecole.objects.create(nom='École recouvrement audit', etat='VALIDE')
        autre = Ecole.objects.create(nom='École étrangère recouvrement', etat='VALIDE')
        self.eleves = []
        for ecole in (self.ecole, autre):
            classe = Classe.objects.create(ecole=ecole, nom='1ère année', niveau='PRIMAIRE_1', annee_scolaire='2026-2027')
            self.eleves.append(Eleve.objects.create(classe=classe, matricule=f'REC-AUD-{ecole.pk}', nom='Bah', prenom='Awa', sexe='F'))
        profil = self.user.profil
        profil.ecole = self.ecole
        profil.role = 'ADMIN'
        profil.actif = profil.is_validated = True
        profil.save()
        self.client.force_login(self.user)
        self.today = timezone.localdate()
        self.data = {'eleve': self.eleves[0].pk, 'montant': 10000, 'date_debut': self.today, 'date_fin': self.today + timedelta(days=30), 'statut': 'ACTIF', 'alerte_avant_jours': 7}

    def test_creation_et_modification_refusent_eleve_autre_ecole(self):
        url = reverse('depenses:ajouter_abonnement_informatique')
        response = self.client.post(url, {**self.data, 'eleve': self.eleves[1].pk})
        self.assertEqual(response.status_code, 200)
        self.assertIn('eleve', response.context['form'].errors)
        self.assertFalse(AbonnementInformatique.objects.exists())
        abonnement = AbonnementInformatique.objects.create(eleve=self.eleves[0], montant=10000, date_debut=self.today, date_fin=self.today + timedelta(days=30))
        response = self.client.post(reverse('depenses:modifier_abonnement_informatique', args=[abonnement.pk]), {**self.data, 'eleve': self.eleves[1].pk})
        self.assertEqual(response.status_code, 200)
        self.assertIn('eleve', response.context['form'].errors)
        abonnement.refresh_from_db()
        self.assertEqual(abonnement.eleve, self.eleves[0])

    def test_montant_negatif_refuse_sans_creation(self):
        response = self.client.post(reverse('depenses:ajouter_abonnement_informatique'), {**self.data, 'montant': -1})
        self.assertEqual(response.status_code, 200)
        self.assertIn('montant', response.context['form'].errors)
        self.assertFalse(AbonnementInformatique.objects.exists())

    def test_filtre_expire_et_export_reconnaissent_la_date_depassee(self):
        abonnement = AbonnementInformatique.objects.create(eleve=self.eleves[0], montant=10000, date_debut=self.today - timedelta(days=31), date_fin=self.today - timedelta(days=1))
        response = self.client.get(reverse('depenses:liste_abonnements_informatique'), {'filtre': 'expire'})
        self.assertEqual([abo.pk for abo in response.context['page_obj']], [abonnement.pk])
        response = self.client.get(reverse('depenses:liste_abonnements_informatique'), {'filtre': 'actif'})
        self.assertEqual(response.context['page_obj'].paginator.count, 0)
        response = self.client.get(reverse('depenses:export_informatique_excel'))
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.cell(2, 7).value, 'Expiré')
        abonnement.refresh_from_db()
        self.assertEqual(abonnement.statut, 'ACTIF')

    def test_delai_alerte_personnalise_et_zero_sont_respectes(self):
        abonnement = AbonnementInformatique(eleve=self.eleves[0], date_fin=self.today + timedelta(days=10), alerte_avant_jours=14)
        self.assertTrue(abonnement.est_proche_expiration)
        abonnement.alerte_avant_jours = 0
        self.assertFalse(abonnement.est_proche_expiration)
        abonnement.date_fin = self.today
        self.assertTrue(abonnement.est_proche_expiration)
        abonnement.statut = 'SUSPENDU'
        self.assertFalse(abonnement.est_proche_expiration)

    def test_dates_invalides_refusees_sur_liste_et_exports(self):
        for route in ('liste_module_simple', 'export_module_simple_excel', 'export_module_simple_pdf'):
            for filtres in ({'date_debut': 'invalide'}, {'date_debut': '2026-12-31', 'date_fin': '2026-01-01'}):
                with self.subTest(route=route, filtres=filtres):
                    self.assertEqual(self.client.get(reverse('depenses:' + route, args=['cuisine']), filtres).status_code, 400)

    def test_export_respecte_le_filtre_de_la_liste(self):
        DepenseCuisine.objects.create(designation='Achat ancien', montant=300, date=self.today - timedelta(days=10), cree_par=self.user)
        DepenseCuisine.objects.create(designation='Achat récent', montant=400, date=self.today, cree_par=self.user)
        response = self.client.get(reverse('depenses:export_module_simple_excel', args=['cuisine']), {'date_debut': self.today.isoformat()})
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 2).value, 'Achat récent')
